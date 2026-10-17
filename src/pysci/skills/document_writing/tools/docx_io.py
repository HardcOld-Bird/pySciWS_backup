"""docx_io —— 用 python-docx 结构化读写 .docx。

读：按文档顺序保留 标题层级 / 正文 / 项目符号 / 表格 → Markdown。
写：从 Block 规格或约定式 Markdown 生成 .docx（标题/正文/项目符号/表格）。

用于「必须从 docx 数据源取信息」或「必须交付 docx」的场景。仅支持 .docx
（OOXML），不支持旧版 .doc。
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

# python-docx 是 [writing] 可选依赖；缺失时给出清晰指引而非裸 ImportError。
try:
    from docx import Document
except ImportError as e:  # pragma: no cover
    raise ImportError("python-docx 未安装。请运行：uv sync --extra writing") from e


# ---------------------------------------------------------------------------
# 数据结构
# ---------------------------------------------------------------------------
@dataclass
class DocxBlock:
    """文档中的一个块（读/写共用）。"""

    kind: str  # "heading" | "paragraph" | "bullet" | "table" | "quote"
    text: str = ""
    level: int = 0  # heading 级别(1..)；bullet 缩进级别(0..)
    rows: list[list[str]] | None = None  # table：行→列


# ---------------------------------------------------------------------------
# 读取
# ---------------------------------------------------------------------------
#: OOXML 命名空间：正文 run 文本（``w:t``）与 OMML 内联公式文本（``m:t``）。
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_TEXT_TAGS = (f"{{{_W_NS}}}t", f"{{{_M_NS}}}t")


def _para_text(p: Any) -> str:
    """段落全文本：按**文档序**遍历 ``w:t`` 与 ``m:t``（内联 OMML 公式文本一并取回）。

    python-docx 的 ``paragraph.text`` 只拼 ``w:r/w:t``，会**静默丢掉** ``<m:oMath>`` 里的
    内联公式文本——正文里的数值/单位常整段是公式（backlog 20261017-052625：实测「4 cm」
    等处变空白，差点据此误判无规格）。故不依赖 ``.text``，自行按序取两类 ``t`` 节点。
    """
    return "".join(
        node.text for node in p.iter() if node.tag in _TEXT_TAGS and node.text
    )


def _heading_level(style_name: str) -> int | None:
    """把 Word 样式名映射为标题级别；非标题返回 None。"""
    name = (style_name or "").strip()
    if name == "Title":
        return 1
    if name.startswith("Heading"):
        tail = name[len("Heading") :].strip()
        try:
            return max(1, int(tail))
        except ValueError:
            return 1
    return None


def _is_bullet(style_name: str) -> bool:
    return (style_name or "").strip().startswith("List Bullet")


def _bullet_level_from_style(style_name: str) -> int:
    tail = (style_name or "").strip()[len("List Bullet") :].strip()
    try:
        return max(0, int(tail) - 1)
    except ValueError:
        return 0


def _cell_text(cell: Any) -> str:
    """单元格全文本（含内联 OMML 公式）：逐段按文档序取 ``w:t``/``m:t``，段间以换行接。"""
    return (
        "\n".join(_para_text(p._p) for p in cell.paragraphs).strip().replace("\n", " ")
    )


def _table_to_rows(table: Any) -> list[list[str]]:
    return [[_cell_text(c) for c in row.cells] for row in table.rows]


def read_docx(path: str | Path) -> list[DocxBlock]:
    """按文档顺序读取 .docx 的块序列。"""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"docx 不存在：{path}")
    if path.suffix.lower() != ".docx":
        raise ValueError(f"仅支持 .docx（不支持旧版 .doc）：{path.name}")

    doc = Document(str(path))
    blocks: list[DocxBlock] = []

    # iter_inner_content 保持段落/表格交错顺序（python-docx >= 1.1）
    items = (
        list(doc.iter_inner_content())
        if hasattr(doc, "iter_inner_content")
        else [
            *doc.paragraphs,
            *doc.tables,
        ]
    )
    for item in items:
        if item.__class__.__name__ == "Table":
            rows = _table_to_rows(item)
            if rows:
                blocks.append(DocxBlock(kind="table", rows=rows))
            continue
        text = _para_text(item._p).strip()
        if not text:
            continue
        style = getattr(item.style, "name", "") or ""
        lvl = _heading_level(style)
        if lvl is not None:
            blocks.append(DocxBlock(kind="heading", text=text, level=lvl))
        elif _is_bullet(style):
            blocks.append(
                DocxBlock(
                    kind="bullet", text=text, level=_bullet_level_from_style(style)
                )
            )
        else:
            blocks.append(DocxBlock(kind="paragraph", text=text))
    return blocks


# ---------------------------------------------------------------------------
# run 级标注抽取（写作自检 / reviewer 机械对照）
# ---------------------------------------------------------------------------
@dataclass
class MarkedRun:
    """一个带格式标注的 run（段落内 0-based 序号）。"""

    para_index: int
    run_index: int
    marks: list[
        str
    ]  # 如 ["strike", "shd=FF0000"]；顺序稳定（strike/shd/highlight/bold/color）
    text: str


def _val_on(el: Any) -> bool:
    """``w:strike`` / ``w:b`` 的开关语义：缺省 val 即为开，0/false/none 为关。"""
    val = el.get(f"{{{_W_NS}}}val")
    return str(val if val is not None else "true").lower() not in (
        "0",
        "false",
        "none",
        "off",
    )


def _run_marks(run: Any) -> list[str]:
    """run 的标注清单（实测惯用面：strike / shd 底色 / bold，另含 highlight / color）。"""
    rpr = run._element.rPr
    if rpr is None:
        return []
    out: list[str] = []
    strike = rpr.find(f"{{{_W_NS}}}strike")
    if strike is not None and _val_on(strike):
        out.append("strike")
    shd = rpr.find(f"{{{_W_NS}}}shd")
    if shd is not None:
        fill = str(shd.get(f"{{{_W_NS}}}fill") or "").strip().upper()
        if fill and fill != "AUTO":
            out.append(f"shd={fill}")
    hl = rpr.find(f"{{{_W_NS}}}highlight")
    if hl is not None:
        val = str(hl.get(f"{{{_W_NS}}}val") or "").strip()
        if val and val != "none":
            out.append(f"highlight={val}")
    bold = rpr.find(f"{{{_W_NS}}}b")
    if bold is not None and _val_on(bold):
        out.append("bold")
    color = rpr.find(f"{{{_W_NS}}}color")
    if color is not None:
        val = str(color.get(f"{{{_W_NS}}}val") or "").strip().upper()
        if val and val != "AUTO":
            out.append(f"color={val}")
    return out


def extract_run_marks(
    path: str | Path, *, para: str = ""
) -> tuple[list[tuple[int, str, list[MarkedRun]]], int]:
    """抽取正文段落的 run 级标注（strike / shd 底色 / highlight / bold / color）。

    Args:
        path: .docx 路径。
        para: 段落文本关键词过滤（子串匹配，空串＝不过滤）。

    Returns:
        ``(grouped, total)``：``grouped`` = ``[(段落序号, 段落全文, [MarkedRun…]), …]``，
        只含**有标注 run** 且文本命中关键词的段落；``total`` = 标注 run 总数。
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"docx 不存在：{path}")
    if path.suffix.lower() != ".docx":
        raise ValueError(f"仅支持 .docx（不支持旧版 .doc）：{path.name}")

    doc = Document(str(path))
    grouped: list[tuple[int, str, list[MarkedRun]]] = []
    total = 0
    for i, p in enumerate(doc.paragraphs):
        text = _para_text(p._p)
        if para and para not in text:
            continue
        marked = []
        for j, run in enumerate(p.runs):
            marks = _run_marks(run)
            if marks:
                marked.append(MarkedRun(i, j, marks, run.text))
        if marked:
            grouped.append((i, text, marked))
            total += len(marked)
    return grouped, total


# ---------------------------------------------------------------------------
# 渲染为 Markdown
# ---------------------------------------------------------------------------
def _rows_to_md_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(norm[0]) + " |", "|" + "|".join(["---"] * width) + "|"]
    for r in norm[1:]:
        out.append("| " + " | ".join(cell.replace("|", "\\|") for cell in r) + " |")
    return "\n".join(out)


def blocks_to_markdown(blocks: list[DocxBlock], *, source_name: str = "") -> str:
    parts: list[str] = []
    if source_name:
        parts.append(f"# DOCX 提取：{source_name}\n")
    for b in blocks:
        if b.kind == "heading":
            parts.append(f"{'#' * min(b.level, 6)} {b.text}\n")
        elif b.kind == "bullet":
            parts.append(f"{'  ' * b.level}- {b.text}\n")
        elif b.kind == "table":
            parts.append(_rows_to_md_table(b.rows or []) + "\n")
        elif b.kind == "quote":
            parts.append(f"> {b.text}\n")
        else:
            parts.append(f"{b.text}\n")
    return "\n".join(parts)


def docx_to_markdown(path: str | Path) -> str:
    """一步到位：读取 .docx → Markdown。"""
    path = Path(path)
    return blocks_to_markdown(read_docx(path), source_name=path.name)


# ---------------------------------------------------------------------------
# 写入
# ---------------------------------------------------------------------------
_BULLET_STYLES = (
    "List Bullet",
    "List Bullet 2",
    "List Bullet 3",
    "List Bullet 4",
    "List Bullet 5",
)


def _add_block(doc: Any, b: DocxBlock) -> None:
    if b.kind == "heading":
        doc.add_heading(b.text, level=max(1, min(b.level, 9)))
    elif b.kind == "bullet":
        doc.add_paragraph(
            b.text, style=_BULLET_STYLES[min(b.level, len(_BULLET_STYLES) - 1)]
        )
    elif b.kind == "table":
        rows = b.rows or []
        if not rows:
            return
        n_c = max(len(r) for r in rows)
        tbl = doc.add_table(rows=len(rows), cols=n_c)
        tbl.style = "Table Grid"
        for r, row in enumerate(rows):
            for c in range(n_c):
                tbl.cell(r, c).text = row[c] if c < len(row) else ""
    elif b.kind == "quote":
        para = doc.add_paragraph()
        para.add_run(b.text).italic = True
    else:
        doc.add_paragraph(b.text)


def build_docx(
    blocks: list[DocxBlock], out_path: str | Path, *, title: str | None = None
) -> Path:
    """从 Block 列表新建 .docx；title 给定时加一个 Title 级标题。"""
    doc = Document()
    if title:
        doc.add_heading(title, level=0)
    for b in blocks:
        _add_block(doc, b)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def add_block_to(path: str | Path, block: DocxBlock) -> Path:
    """向已有 .docx 追加一个块并保存。"""
    path = Path(path)
    doc = Document(str(path))
    _add_block(doc, block)
    doc.save(str(path))
    return path
