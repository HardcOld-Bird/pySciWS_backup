"""字体**全槽**复制（latin+ea+cs+panose）：绕开 python-pptx 的 latin-only 默认。

python-pptx 的 ``run.font.name`` 只写 ``<a:latin>``；中文实际走 ``<a:ea>``。模板的
rPr 通常带 latin/ea/cs 三槽及 ``panose`` 属性——只抄 latin 会让中文回退到主题默认字体，
渲染后字形/折行全变（tb-5 返工 3 轮根因）。本模块提供：

- :func:`font_slots`：读一个 run 的三槽（typeface/panose/pitchFamily/charset）；
- :func:`set_font_slots`：直接写三槽；
- :func:`copy_font_slots`：把模板 run 的三槽**整元素复制**到目标 run；
- :func:`copy_slots_to_shape`：按 run 序把模板文本框的槽位复制到目标文本框。

audit（``slides audit``）用 :func:`font_slots` 校验每标题/胶囊/正文 run 是否含 ``<a:ea>``。
"""

from __future__ import annotations

import copy
from typing import Any

from pptx.oxml.ns import qn

#: rPr 内的字体槽元素（顺序即写入顺序；`a:sym` 不在本工具范围）。
FONT_SLOTS: tuple[str, ...] = ("latin", "ea", "cs")


def _rpr(run: Any) -> Any:
    return run._r.get_or_add_rPr()


def font_slots(run: Any) -> dict[str, dict[str, str]]:
    """读 run 的字体槽：``{latin:{typeface,panose,…}, ea:{…}, cs:{…}}``（缺槽不出现）。"""
    rpr = _rpr(run)
    out: dict[str, dict[str, str]] = {}
    for tag in FONT_SLOTS:
        node = rpr.find(qn(f"a:{tag}"))
        if node is not None:
            out[tag] = {k: v for k, v in node.attrib.items()}
    return out


def set_font_slots(
    run: Any,
    *,
    latin: str,
    ea: str | None = None,
    cs: str | None = None,
    panose: str | None = None,
    pitch_family: str | None = None,
    charset: str | None = None,
) -> None:
    """写入三槽（ea/cs 缺省沿用 latin）；``panose`` 同时写三槽（同模板做法）。"""
    rpr = _rpr(run)
    for tag in FONT_SLOTS:
        old = rpr.find(qn(f"a:{tag}"))
        if old is not None:
            rpr.remove(old)
    for tag, face in (("latin", latin), ("ea", ea or latin), ("cs", cs or latin)):
        node = rpr.makeelement(qn(f"a:{tag}"), {})
        node.set("typeface", face)
        if panose:
            node.set("panose", panose)
        if pitch_family:
            node.set("pitchFamily", str(pitch_family))
        if charset:
            node.set("charset", str(charset))
        rpr.append(node)


def copy_font_slots(template_run: Any, target_run: Any) -> None:
    """把模板 run 的 latin/ea/cs（含 panose 等属性）整元素复制到目标 run。"""
    src = _rpr(template_run)
    dst = _rpr(target_run)
    for tag in FONT_SLOTS:
        old = dst.find(qn(f"a:{tag}"))
        if old is not None:
            dst.remove(old)
        node = src.find(qn(f"a:{tag}"))
        if node is not None:
            dst.append(copy.deepcopy(node))


def copy_slots_to_shape(template_shape: Any, target_shape: Any) -> int:
    """把模板形状文本框的槽位按 run 序复制到目标形状；返回复制的 run 数。"""
    try:
        src_runs = [r for p in template_shape.text_frame.paragraphs for r in p.runs]
        dst_runs = [r for p in target_shape.text_frame.paragraphs for r in p.runs]
    except Exception:  # noqa: BLE001 - 非文本框
        return 0
    for src, dst in zip(src_runs, dst_runs):
        copy_font_slots(src, dst)
    return min(len(src_runs), len(dst_runs))


def slots_complete(slots: dict[str, dict[str, str]]) -> bool:
    """三槽齐全且各有 typeface（audit 的「含 <a:ea>」判据）。"""
    return all(slots.get(t, {}).get("typeface") for t in FONT_SLOTS)
