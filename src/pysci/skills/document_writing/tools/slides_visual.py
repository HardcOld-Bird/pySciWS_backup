"""slides_visual —— PPTX 逐页渲染与视觉验收（机制性视觉检查闸）。

用户硬要求：**不要让组员在不经视觉检查的情况下就提交 ppt 产物**——机制保证优先于自觉。
本模块是那条机制的**生产端与验收端**：

- 生产端 :func:`render_pptx_pages`：固化 pptx→pdf→逐页 PNG 链（LibreOffice + pymupdf），
  产出 `<deck>_render/page-NNN.png`，供组员 Read 做视觉判读、并作为交付**必附**的渲染图。
- 验收端 :func:`verify_pptx_render`：机械校验**渲染存在**、**页数 = 幻灯片数**、**非空白页**；
  任一不过即判 FAIL——把「组员自觉渲染」升级为「缺渲染即打回」。
- 几何端 :func:`geometry_report`（`slides geometry`）：**只读、不渲染**地硬判版面几何硬伤——
  外框边被更高 z 序图片**覆盖**、图元**越界**（画布/外框）、**图片互叠**面积。渲染闸只管
  「渲染出来了没 / 是不是空白」，命中不了「渲染出来了但版面是坏的」；这一端补上后者。
- 预设端（同一命令内）：组会 PPT 预设方案的三项机械检查——标题「·…·」形态、标题框宽随文
  适配（±10%）、非整页图黑描边。方案单源见 document-writing 技能
  ``references/ppt-preset.md``。

设计取向：验收只读**已落盘的渲染图**（不自行渲染），故「缺渲染」是一个可被机械抓住的
真实状态（组员忘了渲染 → 渲染目录空 → FAIL），而非被本命令静默补齐。渲染与验收分成两条
命令（`compose slides render` / `compose slides verify`），前者产附件、后者做闸门。
"""

from __future__ import annotations

import json
import unicodedata
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from . import office_convert, pdf_render, pptx_io

#: 逐页 PNG 的文件名前缀（与 :func:`pdf_render.render_pdf_pages` 默认一致）。
PAGE_PREFIX = "page"

#: 判定「非空白页」的最小墨迹占比：与**背景（众数灰度）异色**的像素占全图比例低于此值
#: 即视为空白页。取值极小（0.2%）——只抓真正的空页（渲染失败/整页纯色），不误伤
#: 内容稀少但有文字/图元的正常页。
DEFAULT_MIN_INK_RATIO = 0.002

#: 背景与墨迹的灰度差阈值（0–255）：差 > 此值才算「墨迹」。
_INK_GRAY_DELTA = 24


def default_render_dir(pptx: str | Path) -> Path:
    """逐页渲染图的**约定目录**：``<pptx 同目录>/<stem>_render/``。

    验收端按同一约定定位，故 `compose slides render`（生产）与 `compose slides verify`
    （验收）不必显式互传路径即可对齐——这是「逐页渲染 PNG 为交付必附件」的落点。
    """
    p = Path(pptx)
    return p.parent / f"{p.stem}_render"


@dataclass
class RenderResult:
    """一次 pptx → pdf → 逐页 PNG 渲染的结果。"""

    pptx: Path
    pdf: Path
    out_dir: Path
    pngs: list[Path] = field(default_factory=list)
    slide_count: int = 0

    def summary(self) -> str:
        return (
            f"[slides render] {self.pptx.name}：{self.slide_count} 页 → 渲染 "
            f"{len(self.pngs)} 张 PNG 于 {self.out_dir}"
        )


@dataclass
class VerifyReport:
    """一次 pptx 视觉验收的结果（渲染存在 / 页数一致 / 非空白）。"""

    pptx: Path
    render_dir: Path
    slide_count: int = 0
    png_count: int = 0
    blank_pages: list[str] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)
    #: 渲染目录的来源（:func:`resolve_render_dir` 的第 2 返回值）：``explicit``/``default``/
    #: ``discovered``/``missing``/``ambiguous``。非 ``default`` 时写进摘要——「用了别的目录」
    #: 这件事必须可见，免得静默换面。
    render_dir_source: str = "default"

    @property
    def ok(self) -> bool:
        """无任何问题即通过。"""
        return not self.problems

    def summary(self) -> str:
        head = "PASS" if self.ok else "FAIL"
        src_note = (
            ""
            if self.render_dir_source in ("default", "")
            else f"，目录来源={self.render_dir_source}"
        )
        lines = [
            f"[slides verify] {head} {self.pptx.name}：幻灯片 {self.slide_count} 页 / "
            f"渲染 {self.png_count} 张（{self.render_dir}{src_note}）"
        ]
        lines.extend(f"    ✗ {p}" for p in self.problems)
        return "\n".join(lines)


def render_pptx_pages(
    pptx: str | Path,
    *,
    out_dir: str | Path | None = None,
    dpi: int = 140,
) -> RenderResult:
    """把 .pptx 渲染为逐页 PNG（pptx → pdf → page-NNN.png）。

    Args:
        pptx: 源 .pptx。
        out_dir: 输出目录；默认 :func:`default_render_dir`（``<stem>_render/``）。中间 PDF
            也落此目录（``<stem>.pdf``），供回溯与在 PDF 阅读器里放大核对。
        dpi: 渲染分辨率。

    Returns:
        RenderResult：含输出目录、逐页 PNG 路径与幻灯片页数。

    Raises:
        FileNotFoundError: pptx 不存在。
        office_convert.LibreOfficeNotInstalled: 未检测到 LibreOffice（渲染链依赖它）。
    """
    src = Path(pptx)
    if not src.exists():
        raise FileNotFoundError(f"pptx 不存在：{src}")
    slide_count = len(pptx_io.read_pptx(src))
    out = Path(out_dir) if out_dir else default_render_dir(src)
    out.mkdir(parents=True, exist_ok=True)
    pdf, _backend = office_convert.convert_ex(src, "pdf", out_dir=out)
    print(f"[slides render] PDF 后端：{_backend}")
    pngs = pdf_render.render_pdf_pages(pdf, out_dir=out, dpi=dpi, prefix=PAGE_PREFIX)
    return RenderResult(
        pptx=src, pdf=pdf, out_dir=out, pngs=pngs, slide_count=slide_count
    )


def page_ink_ratio(png: str | Path) -> float:
    """一页渲染图里「墨迹」（与背景异色）像素占全图的比例。

    背景取灰度直方图的**众数**（最常见灰度），与背景差 > :data:`_INK_GRAY_DELTA` 的像素
    计为墨迹。空白页返回 ~0。
    """
    from PIL import Image

    with Image.open(png) as im:
        hist = im.convert("L").histogram()
    total = sum(hist) or 1
    bg = max(range(256), key=lambda v: hist[v])
    ink = sum(hist[v] for v in range(256) if abs(v - bg) > _INK_GRAY_DELTA)
    return ink / total


def page_is_blank(
    png: str | Path, min_ink_ratio: float = DEFAULT_MIN_INK_RATIO
) -> bool:
    """该页是否**空白**（墨迹占比低于阈值）。"""
    return page_ink_ratio(png) < min_ink_ratio


def _png_count(d: Path) -> int:
    try:
        return len(list(d.glob("*.png")))
    except OSError:
        return 0


def resolve_render_dir(
    pptx: str | Path, *, render_dir: str | Path | None = None, slide_count: int = 0
) -> tuple[Path, str, list[str]]:
    """定位逐页渲染图目录，返回 ``(目录, 来源, 同级候选)``。

    解析顺序（backlog 20261016-042501-devops：不再逼交付维护双份渲染目录）：

    1. **显式** ``--render-dir``（最高优先，原样采用）；
    2. 约定 ``<stem>_render/``（原默认，非空即用——默认行为兼容）；
    3. pptx **同级**含 PNG 的子目录：多个时取页数吻合者；**唯一但张数与幻灯片页数不符
       的旁路目录直接跳过**（返回 ``mismatch``）——否则把 `插图/` 一类素材目录当渲染件，
       会把「缺渲染」误报成「页数不符」（20261019-192536-reviewer 实案：4≠15）。
       仍歧义（多个候选且无一吻合）→ 报候选、判 ``ambiguous``（不猜）。

    来源取值：``explicit`` / ``default`` / ``discovered`` / ``missing`` / ``mismatch`` /
    ``ambiguous``。
    """
    src = Path(pptx)
    if render_dir is not None and str(render_dir).strip():
        return Path(render_dir), "explicit", []
    default = default_render_dir(src)
    if _png_count(default) > 0:
        return default, "default", []
    parent = src.parent
    cands = (
        [d for d in sorted(parent.iterdir()) if d.is_dir() and _png_count(d) > 0]
        if parent.is_dir()
        else []
    )
    if not cands:
        return default, "missing", []
    desc = [str(d) for d in cands]
    exact = [d for d in cands if _png_count(d) == slide_count]
    if len(exact) == 1:
        return exact[0], "discovered", desc
    if len(cands) == 1:
        only = cands[0]
        n = _png_count(only)
        if slide_count and n != slide_count:
            return default, "mismatch", desc
        return only, "discovered", desc
    return default, "ambiguous", desc


def verify_pptx_render(
    pptx: str | Path,
    *,
    render_dir: str | Path | None = None,
    min_ink_ratio: float = DEFAULT_MIN_INK_RATIO,
) -> VerifyReport:
    """机械验收一份 pptx 的**逐页渲染**（渲染存在 / 页数一致 / 非空白）。

    只读已落盘的渲染图（不自行渲染）——故「缺渲染」会被如实抓住。渲染目录由
    :func:`resolve_render_dir` 定位（显式 → 约定 → 同级发现），故渲染件落在用户指定的
    `render/` 一类目录里也能过关，不必再复制一份到 `<stem>_render/`。判据三条任一不过即
    计入 :attr:`VerifyReport.problems`，:attr:`VerifyReport.ok` 随之 False。

    Args:
        pptx: 源 .pptx。
        render_dir: 渲染图目录（显式指定时最高优先）；缺省走 :func:`resolve_render_dir`。
        min_ink_ratio: 空白页判定阈值（墨迹占比低于它即空白）。

    Returns:
        VerifyReport。
    """
    src = Path(pptx)
    problems: list[str] = []
    if not src.exists():
        return VerifyReport(
            pptx=src,
            render_dir=Path(render_dir) if render_dir else default_render_dir(src),
            problems=[f"产物不存在：{src}"],
        )

    slide_count = len(pptx_io.read_pptx(src))
    rdir, how, cands = resolve_render_dir(
        src, render_dir=render_dir, slide_count=slide_count
    )
    pngs = sorted(rdir.glob("*.png")) if rdir.is_dir() else []

    if how == "ambiguous":
        problems.append(
            "渲染目录歧义：pptx 同级有多个含 PNG 的目录（"
            + "、".join(cands)
            + "）——用 `--render-dir` 指定"
        )
    elif how == "mismatch":
        problems.append(
            "已跳过张数与幻灯片页数不一致的旁路渲染目录（"
            + "、".join(cands)
            + f"）；缺 {default_render_dir(src).name}/ 或显式 `--render-dir`——"
            "先跑 `compose slides render` 并在交付中附上"
        )
    elif not pngs:
        problems.append(
            f"缺渲染：未找到逐页渲染 PNG（{rdir}）——先跑 `compose slides render` 并在交付中附上"
        )
    elif len(pngs) != slide_count:
        problems.append(
            f"页数不符：渲染 {len(pngs)} 页 ≠ 幻灯片 {slide_count} 页（目录 {rdir}）"
        )

    blank = [p.name for p in pngs if page_is_blank(p, min_ink_ratio)]
    if blank:
        problems.append(f"空白页：{'、'.join(blank)}")

    return VerifyReport(
        pptx=src,
        render_dir=rdir,
        slide_count=slide_count,
        png_count=len(pngs),
        blank_pages=blank,
        problems=problems,
        render_dir_source=how,
    )


# ---------------------------------------------------------------------------
# 几何体检（render-free、只读）：外框边被覆盖 / 图元越界 / 图片互叠
# ---------------------------------------------------------------------------
# 为什么另立一套：渲染闸（verify）只防「忘了渲染」与「渲染成空白/坏页」，对**版面硬伤**
# （文字/图块压住模板外框、图元越出画布或内容框、两张粘贴截图互相叠住）零覆盖——这些是
# 纯几何事实，从 .pptx 的 bbox + z 序即可硬判，不必也不该让审查员自建脚本。故提供
# `slides geometry`：只读、**不写盘**（受审目录只读时也照跑），逐页输出三类数值。
#
# EMU → 英寸。
_EMU_IN = 914400.0
#: 越界判定容差（英寸）：小于它的超出视为对齐误差，不报。
DEFAULT_TOLERANCE_IN = 0.05
#: 图片互叠的默认面积阈值（平方英寸）：小于它的交叠不报（相邻贴合、圆角边不算）。
DEFAULT_OVERLAP_IN2 = 1.0
#: 每条边被覆盖率的采样点数。
_EDGE_SAMPLES = 400
#: 视作「背景图」的面积占比（≥ 此值不参与互叠判定——满版底图与内容图天然相交）。
_BACKGROUND_AREA_FRAC = 0.8

# ---------------------------------------------------------------------------
# 组会 PPT 预设方案（references/ppt-preset.md）的三项机械判据
# ---------------------------------------------------------------------------
#: 标题区（预设第 3 条）：左 0.60 / 上 0.44 / 高 0.44 in。普通文本框落在该区亦视为标题
#: （用户原件 `20260918.pptx` 用「文本框 1」承载标题，不是占位符）。
TITLE_LEFT_IN = 0.60
TITLE_TOP_IN = 0.44
TITLE_HEIGHT_IN = 0.44
TITLE_ZONE_TOL_IN = 0.06

#: 标题框宽估算的字符宽系数（em，与字号成比例）。标定基准 = `20260918.pptx` 用户原件
#: （20 pt 等线 Light）：「·现状总结·」6 字实测框宽 1.78 in、「·pySci·」7 字 1.39 in；
#: 按 CJK/拉丁分别计解得 CJK 1.05 em/字、拉丁 0.58 em/字——四页用户标题拟合误差 ≤4%。
TITLE_CJK_EM = 1.05
TITLE_LATIN_EM = 0.58
#: 标题 run 未显式写字号时的回退（PowerPoint 文本框默认 18 pt）。
DEFAULT_TITLE_PT = 18.0
#: 标题框宽适配容差（相对估算宽的 ±10%）。
DEFAULT_TITLE_WIDTH_TOL = 0.10
#: 文本框默认左右内边距（英寸）：``<a:bodyPr>`` 未显式 lIns/rIns 时 PowerPoint 默认 10%
#:（文本框宽 0.1 in/边）。估算所需框宽时必须计入，否则「按字宽设框」会因可用宽度不足折行。
DEFAULT_TITLE_INSET_IN = 0.1

#: 整页满幅图豁免阈值：图片面积 ≥ 画布面积的此比例即免描边检查（16:9 满幅 AI 图）。
FULL_PAGE_PIC_FRAC = 0.95

#: 装饰性水印的透明度上限（``alphaModFix/@amt`` 刻度，100000 = 100%）：amt=8000 即 8%。
#: 不透明度不超过此值的图片视为**装饰**（学院正式预设的左侧水印），不参与内容图的
#: 图片互叠 / 描边 / 外框判定——否则水印天然压住内容图与白卡边缘，逐页假阳性。
DECORATIVE_ALPHA_MAX = 10000

#: 预设语义下「首页」页码（内容页标题检查的默认豁免）。
FIRST_PAGE = 1


@dataclass
class ShapeBox:
    """一个形状的包围盒（英寸，左上/右下）+ 页内 z 序（枚举序，越大越靠上）。"""

    z: int
    name: str
    kind: str  # "pic" | "text" | "shape" | "connector"
    x0: float
    y0: float
    x1: float
    y1: float
    rect: bool = False  # 是否为矩形自选图形（外框自动识别用）
    #: 透明装饰图（水印）：不参与内容图几何判定（互叠/描边/外框），见
    #: :func:`_is_decorative_picture`。
    decorative: bool = False

    @property
    def box(self) -> tuple[float, float, float, float]:
        return (self.x0, self.y0, self.x1, self.y1)

    @property
    def width(self) -> float:
        return self.x1 - self.x0

    @property
    def height(self) -> float:
        return self.y1 - self.y0

    @property
    def area(self) -> float:
        return max(0.0, self.width) * max(0.0, self.height)


@dataclass
class EdgeCoverage:
    """外框某条边被**更高 z 序图片**覆盖的比例与区间（区间按 400 点采样还原）。"""

    edge: str
    frac: float
    spans: list[str] = field(default_factory=list)


@dataclass
class SlideGeometry:
    """一页的几何体检结果。"""

    index: int
    layout: str = ""
    boxes: list[ShapeBox] = field(default_factory=list)
    frame: ShapeBox | None = None
    frame_source: str = ""  # "explicit" | "auto" | ""
    #: 自动外框识别存疑时的提示（非违规）：存在近满幅白卡却选中子块，或只选中
    #: 小块且有多个矩形候选。此时 edges/off_frame 不再输出，避免假阳性越界。
    frame_note: str = ""
    edges: list[EdgeCoverage] = field(default_factory=list)
    overlaps: list[str] = field(default_factory=list)
    off_canvas: list[str] = field(default_factory=list)
    off_frame: list[str] = field(default_factory=list)
    #: 连接符 bbox 越界**提示**（非违规）：连接符 bbox 不可信，只提示不计入 violations。
    connector_hints: list[str] = field(default_factory=list)
    #: 缺 ``<a:ext>``（继承尺寸）的形状标签——一票否决项（backlog W1）。
    missing_ext: list[str] = field(default_factory=list)
    #: 预设第 3 条：标题文本不是「·…·」形态。
    title_form_issues: list[str] = field(default_factory=list)
    #: 预设第 3 条：标题框宽超出适配区间（字形宽下限 / 含字距与内边距的上限，±10%）。
    title_width_issues: list[str] = field(default_factory=list)
    #: 预设第 4 条：非整页图缺黑色描边（``<a:ln>`` 缺失或 noFill）。
    picture_border_issues: list[str] = field(default_factory=list)
    #: 透明装饰图（水印）名称——只作读数，不计违规。
    decorative_pics: list[str] = field(default_factory=list)

    @property
    def violations(self) -> int:
        """本页违规项计数（有覆盖的边 + 互叠 + 越界 + 缺 ext + 预设三项）。"""
        return (
            len([e for e in self.edges if e.frac > 0])
            + len(self.overlaps)
            + len(self.off_canvas)
            + len(self.off_frame)
            + len(self.missing_ext)
            + len(self.title_form_issues)
            + len(self.title_width_issues)
            + len(self.picture_border_issues)
        )


@dataclass
class GeometryReport:
    """一份 pptx 的几何体检报告（只读、render-free）。"""

    pptx: Path
    canvas_in: tuple[float, float]
    slides: list[SlideGeometry] = field(default_factory=list)
    #: 跨页标题格式差异（W2，报告级；计入 violations）。
    title_diffs: list[str] = field(default_factory=list)

    @property
    def ext_missing_total(self) -> int:
        """全篇缺 ``<a:ext>`` 的形状总数（W1/W4 风险提示的触发量）。"""
        return sum(len(s.missing_ext) for s in self.slides)

    @property
    def violations(self) -> int:
        return sum(s.violations for s in self.slides) + len(self.title_diffs)

    @property
    def ok(self) -> bool:
        return self.violations == 0

    def summary(self) -> str:
        w, h = self.canvas_in
        lines = [
            f"[slides geometry] {self.pptx.name}：画布 {w:.2f}x{h:.2f} in，{len(self.slides)} 页"
        ]
        for s in self.slides:
            head = f"  第 {s.index} 页"
            if s.frame is not None:
                f = s.frame
                head += (
                    f"：外框 {f.name} ({f.x0:.2f},{f.y0:.2f})-({f.x1:.2f},{f.y1:.2f})"
                    f"（{s.frame_source}）"
                )
            else:
                head += "：未识别外框（用 --frame 指定形状名可启用边覆盖判定）"
            lines.append(head)
            if s.frame_note:
                lines.append(f"      [!] {s.frame_note}")
            for e in s.edges:
                if e.frac > 0:
                    lines.append(
                        f"      {e.edge}: 被覆盖 {e.frac * 100:.1f}%  ->  {e.spans}"
                    )
            for ov in s.overlaps:
                lines.append(f"      图片互叠: {ov}")
            for oc in s.off_canvas:
                lines.append(f"      越出画布: {oc}")
            for of in s.off_frame:
                lines.append(f"      越出外框: {of}")
            for h in s.connector_hints:
                lines.append(f"      提示（非违规）: {h}")
            for me in s.missing_ext:
                lines.append(f"      缺 <a:ext>（继承尺寸，一票否决）: {me}")
            for issue in s.title_form_issues:
                lines.append(f"      标题「·…·」形态: {issue}")
            for issue in s.title_width_issues:
                lines.append(f"      标题框宽失配: {issue}")
            for issue in s.picture_border_issues:
                lines.append(f"      图片描边: {issue}")
            if s.decorative_pics:
                lines.append(
                    f"      装饰水印（豁免互叠/描边/外框判定）: {'、'.join(s.decorative_pics)}"
                )
        for d in self.title_diffs:
            lines.append(f"      标题不一致: {d}")
        if self.ext_missing_total:
            lines.append(
                f"  [!] 风险（W4）：{self.ext_missing_total} 个形状显式尺寸缺失——"
                "**渲染图正常 ≠ WPS 正常**：LibreOffice 会按继承值兜底渲染，WPS 可能按"
                "零/异常尺寸呈现；请补齐 `<a:xfrm><a:ext>` 后复跑。"
            )
        lines.append(
            f"[slides geometry] 合计违规 {self.violations} 项"
            + ("（PASS）" if self.ok else "（需人工判读/返工）")
        )
        return "\n".join(lines)

    def as_dict(self) -> dict[str, Any]:
        w, h = self.canvas_in
        return {
            "pptx": str(self.pptx),
            "canvas_in": [w, h],
            "violations": self.violations,
            "ext_missing_total": self.ext_missing_total,
            "title_diffs": self.title_diffs,
            "slides": [
                {
                    "index": s.index,
                    "layout": s.layout,
                    "frame": (
                        {
                            "name": s.frame.name,
                            "box": [round(v, 3) for v in s.frame.box],
                            "source": s.frame_source,
                        }
                        if s.frame is not None
                        else None
                    ),
                    "frame_note": s.frame_note,
                    "edge_coverage": [
                        {"edge": e.edge, "frac": round(e.frac, 4), "spans": e.spans}
                        for e in s.edges
                        if e.frac > 0
                    ],
                    "picture_overlaps": s.overlaps,
                    "off_canvas": s.off_canvas,
                    "off_frame": s.off_frame,
                    "connector_hints": s.connector_hints,
                    "missing_ext": s.missing_ext,
                    "title_form_issues": s.title_form_issues,
                    "title_width_issues": s.title_width_issues,
                    "picture_border_issues": s.picture_border_issues,
                    "decorative_pics": s.decorative_pics,
                }
                for s in self.slides
            ],
        }

    def as_json(self) -> str:
        return json.dumps(self.as_dict(), ensure_ascii=False, indent=2)


def _shape_kind(shape: Any) -> str:
    tag = getattr(getattr(shape, "_element", None), "tag", "")
    if isinstance(tag, str) and tag.endswith("cxnSp"):
        # 连接符（bentConnector*/straightConnector*）：PowerPoint 存的是**能容纳折线路径**的
        # bbox，路径常远小于它 → 用 bbox 判越界会报出渲染图上不存在的假阳性（backlog
        # 20261016-033419-writing ②）。单列 ``connector``，越界只作提示、不计入 violations。
        return "connector"
    try:
        if int(shape.shape_type) == 13:  # MSO_SHAPE_TYPE.PICTURE
            return "pic"
    except (TypeError, ValueError):
        pass
    if getattr(shape, "has_text_frame", False):
        return "text"
    return "shape"


def _is_rectangle(shape: Any) -> bool:
    """是否为矩形类自选图形（外框自动识别的候选；ROUNDED_RECTANGLE 亦含 RECTANGLE 字样）。"""
    try:
        if int(shape.shape_type) != 1:  # MSO_SHAPE_TYPE.AUTO_SHAPE
            return False
        return "RECTANGLE" in str(shape.auto_shape_type)
    except (AttributeError, TypeError, ValueError):
        return False


def _is_decorative_picture(shape: Any) -> bool:
    """图片是否为**透明装饰**（水印）：blip 上的 ``alphaModFix/@amt``（或 ``alpha/@val``）
    不透明度 ≤ :data:`DECORATIVE_ALPHA_MAX`（默认 10%）。

    学院正式预设的左侧水印是 8% 透明图（``amt=8000``）且带 ``srcRect`` 裁切、按设计不描边、
    天然压住内容图与白卡边缘——它不是内容图，不能进图片互叠/描边/外框判据。判据读 XML
    真值（不透明度只在 blip 上表达，python-pptx 无公开属性）。
    """
    el = getattr(shape, "_element", None)
    if el is None:
        return False
    blip = el.find(f".//{_A_BLIP}")
    if blip is None:
        return False
    for tag, attr in ((_A_ALPHA_MOD_FIX, "amt"), (_A_ALPHA, "val")):
        node = blip.find(tag)
        raw = node.get(attr) if node is not None else None
        if raw is None:
            continue
        try:
            amt = int(raw)
        except (TypeError, ValueError):
            return False
        return 0 <= amt <= DECORATIVE_ALPHA_MAX
    return False


def _shape_box(shape: Any, z: int) -> ShapeBox | None:
    try:
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
    except Exception:  # noqa: BLE001 - 个别形状无几何属性，跳过
        return None
    if None in (left, top, width, height):
        return None
    x0, y0 = left / _EMU_IN, top / _EMU_IN
    kind = _shape_kind(shape)
    return ShapeBox(
        z=z,
        name=getattr(shape, "name", "") or f"#{z}",
        kind=kind,
        x0=x0,
        y0=y0,
        x1=x0 + width / _EMU_IN,
        y1=y0 + height / _EMU_IN,
        rect=_is_rectangle(shape),
        decorative=kind == "pic" and _is_decorative_picture(shape),
    )


def _collect(slide: Any) -> list[ShapeBox]:
    """页内所有**有几何**的形状（z 序取枚举序，含被跳过的无几何形状，保持相对次序）。"""
    boxes: list[ShapeBox] = []
    for z, shape in enumerate(slide.shapes, 1):
        b = _shape_box(shape, z)
        if b is not None:
            boxes.append(b)
    return boxes


#: DrawingML 命名空间与「显式尺寸」元素 ``<a:xfrm><a:ext>``。
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_A_XFRM = f"{{{_A_NS}}}xfrm"
_A_EXT = f"{{{_A_NS}}}ext"
_A_BLIP = f"{{{_A_NS}}}blip"
_A_ALPHA_MOD_FIX = f"{{{_A_NS}}}alphaModFix"
_A_ALPHA = f"{{{_A_NS}}}alpha"


def _has_explicit_ext(shape: Any) -> bool:
    """形状 XML 是否**显式**写了 ``<a:xfrm><a:ext>``（尺寸真值）。

    python-pptx 对占位符会返回**继承**来的 ``left/top/width/height``，读属性**看不出**
    「缺 ext」——LibreOffice 兜底渲染正常，WPS 可能按零/异常尺寸呈现（双盲缺陷）。故判据
    必须读 XML 真值（backlog tb-20261017-044633-4-1016ppt-1 W1）。
    """
    el = getattr(shape, "_element", None)
    if el is None:
        return True  # 非 OOXML 形状：无从检查，不误报
    xfrm = el.find(f".//{_A_XFRM}")
    if xfrm is None:
        return False  # 连 xfrm 都没有 → 尺寸全靠继承（最危险）
    return xfrm.find(_A_EXT) is not None


def _shape_label(shape: Any) -> str:
    """形状的可读标签（名＋文本摘要），供缺 ext 报告定位到页内具体形状。"""
    name = getattr(shape, "name", "") or "?"
    txt = ""
    try:
        if getattr(shape, "has_text_frame", False):
            txt = (shape.text_frame.text or "").strip().replace("\n", " ")
    except Exception:  # noqa: BLE001
        txt = ""
    return f"{name}（{txt[:20]}）" if txt else name


def _missing_ext(slide: Any) -> list[str]:
    """本页所有**缺 ``<a:ext>``** 的形状标签（W1 一票否决项的原料）。"""
    return [_shape_label(sh) for sh in slide.shapes if not _has_explicit_ext(sh)]


#: 标题占位符类型（python-pptx 不可用时退化为按名字识别）。
def _title_placeholder_types() -> set:
    try:
        from pptx.enum.shapes import PP_PLACEHOLDER

        return {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
    except Exception:  # noqa: BLE001
        return set()


def _is_title_shape(shape: Any) -> bool:
    """是否页标题（占位符类型为 TITLE/CENTER_TITLE，或名字含 Title/标题）。"""
    if getattr(shape, "is_placeholder", False):
        try:
            if shape.placeholder_format.type in _title_placeholder_types():
                return True
        except Exception:  # noqa: BLE001
            pass
    nm = getattr(shape, "name", "") or ""
    return nm.startswith("Title") or "标题" in nm


def _title_shape(slide: Any) -> Any | None:
    for sh in slide.shapes:
        if _is_title_shape(sh):
            return sh
    return None


def _title_fmt(shape: Any) -> tuple:
    """标题的**格式指纹**：``(cy,x,y, {(字号,字体)…})``（W2 跨页一致性比对的原子）。

    **不含框宽**：预设第 3 条要求框宽随标题文本自适应，逐页宽度不同正是设计要求——
    把 cx 放进指纹会让「合规的自适应」被判成不一致（backlog 20261018-085835-writing）。

    字号/字体取 **set 去重**（backlog 20261017-101130）：同一格式被拆成多个 run（如
    「·」+「pySci·」、两 run 的 rPr 逐字相同）时**不得因 run 数不同**判不一致——视觉等价
    即一致；只有格式集合本身不同（如混排字号）才算差异。
    """
    el = getattr(shape, "_element", None)
    h = x = y = None
    if el is not None:
        xfrm = el.find(f".//{_A_XFRM}")
        if xfrm is not None:
            ext = xfrm.find(_A_EXT)
            off = xfrm.find(f"{{{_A_NS}}}off")
            if ext is not None:
                h = ext.get("cy")
            if off is not None:
                x, y = off.get("x"), off.get("y")
    styles: set[tuple[float | None, str | None]] = set()
    try:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                styles.add((run.font.size.pt if run.font.size else None, run.font.name))
    except Exception:  # noqa: BLE001
        pass
    return (h, x, y, frozenset(styles))


def _fmt_diff(a: tuple, b: tuple) -> str:
    """两个标题格式指纹的差异项（人读）。"""
    labels = ("高", "左", "上", "字号/字体")
    hits = [lab for lab, av, bv in zip(labels, a, b) if av != bv]
    return "、".join(hits) if hits else "格式"


def _title_inconsistencies(
    titles: list[tuple[int, tuple | None]],
) -> list[str]:
    """跨页标题一致性（W2）：基准＝**出现最多的标题格式**（并列取最早页），其余页差异即报。"""
    keys = [k for _, k in titles if k is not None]
    if not keys:
        return []
    tally: dict[tuple, list[int]] = {}
    for page, k in titles:
        if k is not None:
            tally.setdefault(k, []).append(page)
    base_key = max(tally, key=lambda k: (len(tally[k]), -min(tally[k])))
    base_page = min(tally[base_key])
    return [
        f"第 {p} 页标题格式与基准（第 {base_page} 页）不一致：{_fmt_diff(k, base_key)}"
        for p, k in titles
        if k is not None and k != base_key
    ]


# ---------------------------------------------------------------------------
# 预设方案三项检查（标题「·…·」/ 标题框宽适配 / 非整页图描边）
# ---------------------------------------------------------------------------
_P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_P_SPPR = f"{{{_P_NS}}}spPr"
_P_TXBODY = f"{{{_P_NS}}}txBody"
_A_BODYPR = f"{{{_A_NS}}}bodyPr"
_A_RPR = f"{{{_A_NS}}}rPr"
_A_LN = f"{{{_A_NS}}}ln"
_A_NO_FILL = f"{{{_A_NS}}}noFill"


def _preset_title_shape(slide: Any) -> Any | None:
    """定位内容页标题：占位符/命名标题优先，其次落在预设标题区的文本框。

    用户原件 `20260918.pptx` 的标题是普通文本框（名「文本框 1」），不在标题占位符里；
    预设第 3 条用**位置**把它定义清楚：左 0.60 / 上 0.44 / 高 0.44 in（±0.06）。
    """
    t = _title_shape(slide)
    if t is not None:
        return t
    for sh in slide.shapes:
        if not getattr(sh, "has_text_frame", False):
            continue
        try:
            left = sh.left / _EMU_IN
            top = sh.top / _EMU_IN
            height = sh.height / _EMU_IN
        except (AttributeError, TypeError):
            continue
        if None in (left, top, height):
            continue
        if (
            abs(top - TITLE_TOP_IN) <= TITLE_ZONE_TOL_IN
            and abs(height - TITLE_HEIGHT_IN) <= TITLE_ZONE_TOL_IN
            and abs(left - TITLE_LEFT_IN) <= TITLE_ZONE_TOL_IN
        ):
            return sh
    return None


def _title_text(shape: Any) -> str:
    try:
        return " ".join((shape.text_frame.text or "").split())
    except Exception:  # noqa: BLE001
        return ""


def _has_preset_title_form(text: str) -> bool:
    """标题文本是否为「·标题·」（前后各一「·」，中间非空）。"""
    t = text.strip()
    return (
        len(t) >= 3 and t.startswith("·") and t.endswith("·") and t[1:-1].strip() != ""
    )


def _char_em(ch: str) -> float:
    """单字符的宽度（em）：CJK/宽/歧义字符按 CJK 系数，其余按拉丁系数。"""
    return (
        TITLE_CJK_EM
        if unicodedata.east_asian_width(ch) in ("W", "F", "A")
        else TITLE_LATIN_EM
    )


def _glyph_width_in(shape: Any, *, default_pt: float = DEFAULT_TITLE_PT) -> float:
    """标题**字形宽**（英寸）：按字符宽系数逐字累加，不含字距与内边距（单行取最大段）。

    系数以 `20260918.pptx` 用户原件标定（20 pt 等线 Light，``2026-10-18``）：
    「·现状总结·」6 字实测框宽 1.78 in、「·pySci·」7 字 1.39 in，按 CJK/拉丁分别计
    得 CJK 1.05 em/字、拉丁 0.58 em/字；四页用户标题拟合误差 ≤4%。合规检查的 ±10%
    区间以此为用户认可基准（框宽下限）。
    """
    try:
        paragraphs = list(shape.text_frame.paragraphs)
    except Exception:  # noqa: BLE001
        return 0.0
    widths: list[float] = []
    for para in paragraphs:
        try:
            para_pt = para.font.size.pt if para.font.size else None
        except Exception:  # noqa: BLE001
            para_pt = None
        total = 0.0
        for run in para.runs:
            try:
                pt = run.font.size.pt if run.font.size else (para_pt or default_pt)
            except Exception:  # noqa: BLE001
                pt = para_pt or default_pt
            for ch in run.text:
                if ch in "\r\n":
                    continue
                total += _char_em(ch) * pt / 72.0
        widths.append(total)
    return max(widths) if widths else 0.0


def _run_spacing_in(run: Any, text: str) -> float:
    """一个 run 的字距总宽（英寸）：``a:rPr/@spc`` 单位为 1/100 pt，按每字符一个间隔保守估算。"""
    r_pr = getattr(run, "_r", None)
    r_pr = r_pr.find(_A_RPR) if r_pr is not None else None
    raw = r_pr.get("spc") if r_pr is not None else None
    if not raw:
        return 0.0
    try:
        spc_pt = int(raw) / 100.0
    except (TypeError, ValueError):
        return 0.0
    return len([c for c in text if c not in "\r\n"]) * spc_pt / 72.0


def _title_insets_in(shape: Any) -> tuple[float, float]:
    """标题文本框的左右内边距（英寸）：读 ``bodyPr/@lIns·rIns``（EMU），缺省各 0.1 in。"""
    el = getattr(shape, "_element", None)
    body = el.find(f".//{_P_TXBODY}") if el is not None else None
    body_pr = body.find(_A_BODYPR) if body is not None else None

    def _one(attr: str) -> float:
        raw = body_pr.get(attr) if body_pr is not None else None
        if raw is None:
            return DEFAULT_TITLE_INSET_IN
        try:
            return int(raw) / _EMU_IN
        except (TypeError, ValueError):
            return DEFAULT_TITLE_INSET_IN

    return _one("lIns"), _one("rIns")


def title_text_width_in(
    shape: Any,
    *,
    default_pt: float = DEFAULT_TITLE_PT,
    include_spacing: bool = True,
) -> float:
    """标题的**文字宽**（英寸，单行）：字形宽 + 字距（可选），不含文本框内边距。

    Args:
        shape: 标题形状（文本框或占位符）。
        default_pt: run/段落均未显式写字号时的回退字号。
        include_spacing: 是否计入 ``a:rPr/@spc`` 字距。

    Returns:
        文字宽（英寸）；无文本或读取失败返回 0.0。
    """
    base = _glyph_width_in(shape, default_pt=default_pt)
    if not include_spacing or base <= 0:
        return base
    try:
        paragraphs = list(shape.text_frame.paragraphs)
    except Exception:  # noqa: BLE001
        return base
    widths: list[float] = []
    for para in paragraphs:
        total = 0.0
        for run in para.runs:
            total += _run_spacing_in(run, run.text)
        widths.append(total)
    return base + (max(widths) if widths else 0.0)


def estimate_title_width_in(
    shape: Any,
    *,
    default_pt: float = DEFAULT_TITLE_PT,
    include_spacing: bool = True,
    include_insets: bool = True,
    inset_in: float | None = None,
) -> float:
    """估算标题**所需框宽**（英寸，单行、不折行）：文字宽 + 左右内边距。

    文字宽 = 字形宽（:func:`_glyph_width_in`）+ 字距（:func:`_run_spacing_in`）；内边距读
    ``bodyPr/@lIns·rIns``（缺省各 0.1 in）。写框时按 ``estimate * 1.1`` 设宽即可保证在
    3 pt 字距等场景下不折行（backlog 20261018-085835-writing：旧实现只估字形宽，按它
    +10% 设框会被默认内边距吃掉可用宽度而折行）。

    **合规检查不用本函数**：它比用户认可的字形宽基准多出内边距与字距；框宽适配区间见
    :func:`_preset_title_issues`（下限=字形宽基准，上限=本估算 +10%）。

    Args:
        shape: 标题形状（文本框或占位符）。
        default_pt: run/段落均未显式写字号时的回退字号。
        include_spacing: 是否计入字距。
        include_insets: 是否计入左右内边距。
        inset_in: 单侧内边距（英寸）；``None`` 时读形状自身，读不到用 0.1。

    Returns:
        所需框宽（英寸）；无文本或读取失败返回 0.0。
    """
    width = title_text_width_in(
        shape, default_pt=default_pt, include_spacing=include_spacing
    )
    if not include_insets or width <= 0:
        return width
    left, right = _title_insets_in(shape)
    if inset_in is not None:
        left = right = inset_in
    return width + left + right


def _preset_title_issues(
    shape: Any, *, width_tol: float
) -> tuple[list[str], list[str]]:
    """一页标题的两项机械检查：返回 ``(形态问题, 框宽问题)``。

    框宽适配是**区间**判据：下限 = 字形宽基准 ``(1 - tol)``（不得窄于用户认可的字宽框），
    上限 = 含字距与内边距的所需框宽 ``(1 + tol)``（不得过长）——这样既不误伤随文本自适应
    的窄标题，也不放过「全篇统一 6.01 in」的返工事故。
    """
    text = _title_text(shape)
    if not text:
        return [], []
    form: list[str] = []
    width: list[str] = []
    if not _has_preset_title_form(text):
        form.append(f"标题缺「·…·」形态：{text!r}（预设第 3 条）")
    try:
        actual = shape.width / _EMU_IN
    except (AttributeError, TypeError):
        actual = None
    glyph = _glyph_width_in(shape)
    padded = estimate_title_width_in(shape)
    if actual is not None and glyph > 0:
        lower = glyph * (1 - width_tol)
        upper = max(glyph, padded) * (1 + width_tol)
        if actual < lower or actual > upper:
            width.append(
                f"框宽 {actual:.2f} in 超出适配区间 [{lower:.2f}, {upper:.2f}] in"
                f"（字形估算 {glyph:.2f}、含字距与内边距 {padded:.2f}，"
                f"容差 ±{width_tol * 100:.0f}%；标题 {text!r}）"
            )
    return form, width


def _picture_border_issue(shape: Any) -> str | None:
    """非整页图的描边检查：``<a:ln>`` 缺失或显式 noFill 即报（有描边不究粗细）。"""
    el = getattr(shape, "_element", None)
    if el is None:
        return None
    sp_pr = el.find(_P_SPPR)
    ln = sp_pr.find(_A_LN) if sp_pr is not None else None
    if ln is None:
        return "无描边（缺 <a:ln>；预设第 4 条要求非整页图加黑色描边）"
    if ln.find(_A_NO_FILL) is not None:
        return "描边为 noFill（预设第 4 条要求非整页图加黑色描边）"
    return None


def _detect_frame(boxes: list[ShapeBox], cw: float, ch: float) -> ShapeBox | None:
    """自动挑外框：画布面积的 10%–99.5% 之间、面积最大的矩形。

    上限原为 95%，学院预设的近满幅白卡（画布 96%）因此被排除、回退到焦点块面板并
    制造「其余列越出外框」假阳性（backlog 20261019-050008）。99.5% 只排除真正的
    full-bleed 背景（≈100%），保留白卡。启发式，非稳准——命名外框（``--frame``）
    仍是首选；识别存疑时由 :func:`geometry_report` 给 ``frame_note`` 并抑制越界行。
    """
    canvas_area = cw * ch
    cands = [
        b
        for b in boxes
        if b.rect and 0.10 * canvas_area <= b.area <= 0.995 * canvas_area
    ]
    return max(cands, key=lambda b: b.area) if cands else None


def _edge_coverage(frame: ShapeBox, blockers: list[ShapeBox]) -> list[EdgeCoverage]:
    """外框四条边按点采样，统计被 `blockers`（更高 z 序图片）覆盖的比例与区间。"""
    n = _EDGE_SAMPLES
    edges = {
        "上边": [(frame.x0 + frame.width * k / n, frame.y0) for k in range(n + 1)],
        "下边": [(frame.x0 + frame.width * k / n, frame.y1) for k in range(n + 1)],
        "左边": [(frame.x0, frame.y0 + frame.height * k / n) for k in range(n + 1)],
        "右边": [(frame.x1, frame.y0 + frame.height * k / n) for k in range(n + 1)],
    }
    out: list[EdgeCoverage] = []
    for ename, pts in edges.items():
        hit = [
            any(b.x0 <= px <= b.x1 and b.y0 <= py <= b.y1 for b in blockers)
            for px, py in pts
        ]
        frac = sum(hit) / len(hit)
        runs: list[tuple[int, int]] = []
        start: int | None = None
        for k, h in enumerate(hit):
            if h and start is None:
                start = k
            elif not h and start is not None:
                runs.append((start, k - 1))
                start = None
        if start is not None:
            runs.append((start, len(hit) - 1))
        spans = []
        for a, b in runs:
            if ename in ("上边", "下边"):
                spans.append(f"x {pts[a][0]:.2f}–{pts[b][0]:.2f}")
            else:
                spans.append(f"y {pts[a][1]:.2f}–{pts[b][1]:.2f}")
        out.append(EdgeCoverage(edge=ename, frac=frac, spans=spans))
    return out


def _picture_overlaps(
    pics: list[ShapeBox], cw: float, ch: float, threshold: float
) -> list[str]:
    """图片两两互叠（交叠面积 > ``threshold``）；满版底图（≥80% 画布）不参与。"""
    bg_area = _BACKGROUND_AREA_FRAC * cw * ch
    res: list[str] = []
    for i in range(len(pics)):
        for j in range(i + 1, len(pics)):
            a, b = pics[i], pics[j]
            if a.area >= bg_area or b.area >= bg_area:
                continue
            ox = min(a.x1, b.x1) - max(a.x0, b.x0)
            oy = min(a.y1, b.y1) - max(a.y0, b.y0)
            if ox > 0 and oy > 0 and ox * oy > threshold:
                res.append(
                    f"{a.name} ∩ {b.name} = {ox:.2f}x{oy:.2f} in ({ox * oy:.2f} in²)"
                    f"  区间 x {max(a.x0, b.x0):.2f}–{min(a.x1, b.x1):.2f}"
                    f" y {max(a.y0, b.y0):.2f}–{min(a.y1, b.y1):.2f}"
                )
    return res


def _off_bounds(
    boxes: list[ShapeBox],
    x0: float,
    y0: float,
    x1: float,
    y1: float,
    tol: float,
    what: str,
) -> list[str]:
    """盒越出参照框（画布或外框）超过 ``tol`` 的清单。"""
    res: list[str] = []
    for b in boxes:
        sides: list[str] = []
        if b.x0 < x0 - tol:
            sides.append(f"左 {x0 - b.x0:.2f}in")
        if b.x1 > x1 + tol:
            sides.append(f"右 {b.x1 - x1:.2f}in")
        if b.y0 < y0 - tol:
            sides.append(f"上 {y0 - b.y0:.2f}in")
        if b.y1 > y1 + tol:
            sides.append(f"下 {b.y1 - y1:.2f}in")
        if sides:
            res.append(
                f"{b.name}({b.x0:.2f},{b.y0:.2f},{b.x1:.2f},{b.y1:.2f})"
                f" 越出{what}：{'、'.join(sides)}"
            )
    return res


def geometry_report(
    pptx: str | Path,
    *,
    frame: str | None = None,
    tolerance: float = DEFAULT_TOLERANCE_IN,
    overlap_in2: float = DEFAULT_OVERLAP_IN2,
    title_exempt_pages: Iterable[int] | None = None,
    title_width_tol: float = DEFAULT_TITLE_WIDTH_TOL,
) -> GeometryReport:
    """对 .pptx 做**只读几何体检**：外框边被覆盖率 / 图元越界（画布与外框）/ 图片互叠面积，
    外加组会 PPT 预设方案的三项机械检查。

    纯 bbox + z 序硬判，不渲染、不写盘——受审目录只读时也照跑（渲染闸 :func:`verify_pptx_render`
    才需要可写目录产 PNG）。**连接符**（``bentConnector*``/``straightConnector*``）的越界因 bbox
    不可信而只作**提示**（:attr:`SlideGeometry.connector_hints`），不计入 ``violations``。

    预设三项（见 document-writing ``references/ppt-preset.md``）：

    - **标题「·…·」形态**：内容页标题（占位符/命名标题/预设标题区的文本框）须匹配
      ``·…·``；首页默认豁免，无标题形状的整页复制页自然跳过；
    - **标题框宽适配**：区间判据——下限为字形宽基准 ``(1 - title_width_tol)``，上限为
      含字距与内边距的所需框宽（:func:`estimate_title_width_in`）``(1 + title_width_tol)``；
      实际框宽出区间即报——不得把全篇标题框统一成同一宽度，也不得过窄折行；
    - **非整页图黑描边**：图片缺 ``<a:ln>``（或显式 noFill）即报；面积 ≥ 画布 95% 的整页
      满幅图豁免。

    **透明装饰图**（水印，blip 不透明度 ≤ :data:`DECORATIVE_ALPHA_MAX`，如学院预设的 8%
    左侧水印）不是内容图：不参与图片互叠、描边与外框判定，只在
    :attr:`SlideGeometry.decorative_pics` 里作读数列出（不计违规）。

    Args:
        pptx: 源 .pptx。
        frame: 外框形状名（``--frame``）。缺省自动挑「面积最大的矩形」（10%–99.5%
            画布面积，近满幅白卡也纳入、≈100% 全幅背景排除）；选中子块存疑时给
            :attr:`SlideGeometry.frame_note` 并抑制越界行。
        tolerance: 越界容差（英寸）。
        overlap_in2: 图片互叠面积阈值（平方英寸）。
        title_exempt_pages: 免做标题检查的页码（1-based）；``None`` 时仅豁免首页。
        title_width_tol: 标题框宽容差（相对估算宽的百分比，0.10 = ±10%）。

    Returns:
        GeometryReport。

    Raises:
        FileNotFoundError: pptx 不存在。
    """
    from pptx import Presentation

    src = Path(pptx)
    if not src.exists():
        raise FileNotFoundError(f"pptx 不存在：{src}")
    prs = Presentation(str(src))
    cw, ch = prs.slide_width / _EMU_IN, prs.slide_height / _EMU_IN
    exempt = {FIRST_PAGE} if title_exempt_pages is None else set(title_exempt_pages)

    slides: list[SlideGeometry] = []
    titles: list[tuple[int, tuple | None]] = []
    for i, slide in enumerate(prs.slides, 1):
        boxes = _collect(slide)
        try:
            layout = slide.slide_layout.name or ""
        except Exception:  # noqa: BLE001
            layout = ""
        fr: ShapeBox | None = None
        source = ""
        if frame:
            fr = next((b for b in boxes if b.name == frame), None)
            source = "explicit" if fr is not None else ""
        else:
            fr = _detect_frame(boxes, cw, ch)
            source = "auto" if fr is not None else ""
        frame_note = ""
        if fr is not None and source == "auto":
            canvas_area = cw * ch
            rect_cands = [
                b
                for b in boxes
                if b.rect and 0.10 * canvas_area <= b.area <= 0.995 * canvas_area
            ]
            near_full = [b for b in rect_cands if b.area >= 0.95 * canvas_area]
            frac = fr.area / max(1.0, canvas_area)
            if near_full and frac < 0.95:
                best = max(near_full, key=lambda b: b.area)
                frame_note = (
                    f"外框识别存疑：存在近满幅矩形 {best.name}"
                    f"（{best.area / canvas_area:.0%} 画布），auto 选中 {fr.name}"
                    f"（{frac:.0%}）；建议 --frame 指定"
                )
            elif frac < 0.5 and len(rect_cands) >= 2:
                frame_note = (
                    f"外框识别存疑：auto 选中 {fr.name}（{frac:.0%} 画布）"
                    "且有多个矩形候选；建议 --frame 指定"
                )

        pics = [b for b in boxes if b.kind == "pic" and not b.decorative]
        decorative_pics = [b.name for b in boxes if b.kind == "pic" and b.decorative]
        # 连接符单列：bbox 不可信，越界只作提示，不混进 violations（否则「零容忍」验收会把
        # 渲染图上并不存在的越界当成真缺陷，逼执行端去改本不该改的源形状）。
        judged = [b for b in boxes if b.kind != "connector"]
        connectors = [b for b in boxes if b.kind == "connector"]

        title_form: list[str] = []
        title_width: list[str] = []
        if i not in exempt:
            title = _preset_title_shape(slide)
            if title is not None:
                title_form, title_width = _preset_title_issues(
                    title, width_tol=title_width_tol
                )

        border: list[str] = []
        full_page_area = FULL_PAGE_PIC_FRAC * cw * ch
        for sh in slide.shapes:
            if _shape_kind(sh) != "pic" or _is_decorative_picture(sh):
                continue
            try:
                area = (sh.width or 0) * (sh.height or 0) / (_EMU_IN**2)
            except TypeError:
                continue
            if area >= full_page_area:
                continue
            issue = _picture_border_issue(sh)
            if issue:
                label = getattr(sh, "name", "") or "?"
                border.append(f"{label}：{issue}")

        sg = SlideGeometry(
            index=i,
            layout=layout,
            boxes=boxes,
            frame=fr,
            frame_source=source,
            frame_note=frame_note,
            off_canvas=_off_bounds(judged, 0.0, 0.0, cw, ch, tolerance, "画布"),
            overlaps=_picture_overlaps(pics, cw, ch, overlap_in2),
            connector_hints=[
                f"{s}（连接符 bbox 提示，非违规）"
                for s in _off_bounds(connectors, 0.0, 0.0, cw, ch, tolerance, "画布")
            ],
            missing_ext=_missing_ext(slide),
            title_form_issues=title_form,
            title_width_issues=title_width,
            picture_border_issues=border,
            decorative_pics=decorative_pics,
        )
        if fr is not None and not frame_note:
            sg.edges = _edge_coverage(fr, [b for b in pics if b.z > fr.z])
            sg.off_frame = _off_bounds(
                pics, fr.x0, fr.y0, fr.x1, fr.y1, tolerance, "外框"
            )
        slides.append(sg)
        t = _title_shape(slide)
        titles.append((i, _title_fmt(t) if t is not None else None))

    return GeometryReport(
        pptx=src,
        canvas_in=(cw, ch),
        slides=slides,
        title_diffs=_title_inconsistencies(titles),
    )
