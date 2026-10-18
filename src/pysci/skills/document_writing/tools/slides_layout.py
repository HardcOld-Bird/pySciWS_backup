"""``compose slides layout``：学院预设的逐页排版机械判据（规则为主、机械为辅）。

按 a2 §二.4 的裁决，先实现**成本低、判据硬**的两条（①焦点块轮廓/②同级间距需块分组与
递归切分语义，当前 deck 未暴露组结构，硬判误报率高，暂缓）：

- **焦点块标题对齐**：同页各焦点块标题（胶囊/标题框）上缘对齐或左缘对齐（容差内）。
- **分焦点块引用框**：每框左缘与其焦点块左缘对齐；文献 ≤3 条时一行一文献；多框上缘
  统一、下缘对齐（下缘到页面灰边的绝对距离需白卡几何，暂只判框间一致）。

形状以 **ID** 指定（``--title`` / ``--refs REF:ANCHOR``），判据只读 bbox 与文本，不推断
版式意图——没给 ID 的项直接跳过并注明，不臆造。
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

_EMU_IN = 914400.0
DEFAULT_TOL_IN = 0.05
_REF_MARK_RE = re.compile(r"\[\d+\]")


@dataclass
class LayoutCheck:
    """单项判据读数。"""

    name: str
    ok: bool
    detail: str


@dataclass
class LayoutReport:
    """一页的排版读数（只读）。"""

    pptx: Path
    page: int
    checks: list[LayoutCheck] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)

    @property
    def violations(self) -> int:
        return sum(1 for c in self.checks if not c.ok)

    @property
    def ok(self) -> bool:
        return self.violations == 0

    def summary(self) -> str:
        lines = [
            f"[slides layout] {self.pptx.name} 第 {self.page} 页："
            + ("PASS" if self.ok else f"{self.violations} 项不通过")
        ]
        for c in self.checks:
            lines.append(f"  [{'√' if c.ok else '✗'}] {c.name}：{c.detail}")
        for s in self.skipped:
            lines.append(f"  [—] {s}（跳过）")
        return "\n".join(lines)

    def as_dict(self) -> dict[str, Any]:
        return {
            "pptx": str(self.pptx),
            "page": self.page,
            "violations": self.violations,
            "checks": [
                {"name": c.name, "ok": c.ok, "detail": c.detail} for c in self.checks
            ],
            "skipped": self.skipped,
        }

    def as_json(self) -> str:
        return json.dumps(self.as_dict(), ensure_ascii=False, indent=2)


def _parse_ids(spec: str | None) -> list[int]:
    out: list[int] = []
    for piece in str(spec or "").split(","):
        piece = piece.strip()
        if not piece:
            continue
        try:
            sid = int(piece)
        except ValueError as exc:
            raise ValueError(f"形状 ID 非整数：{piece!r}") from exc
        if sid not in out:
            out.append(sid)
    return out


def parse_ref_pairs(specs: list[str] | None) -> list[tuple[int, int]]:
    """``REF:ANCHOR`` 列表 → ``[(ref_id, anchor_id)]``。"""
    pairs: list[tuple[int, int]] = []
    for raw in specs or []:
        head, _, tail = str(raw).partition(":")
        try:
            ref_id, anchor_id = int(head.strip()), int(tail.strip())
        except ValueError as exc:
            raise ValueError(f"--refs 须为 REF:ANCHOR 形状 ID：{raw!r}") from exc
        pairs.append((ref_id, anchor_id))
    return pairs


def _box(shape: Any) -> tuple[float, float, float, float]:
    return (
        shape.left / _EMU_IN,
        shape.top / _EMU_IN,
        (shape.left + shape.width) / _EMU_IN,
        (shape.top + shape.height) / _EMU_IN,
    )


def _text(shape: Any) -> str:
    try:
        return shape.text_frame.text or ""
    except Exception:  # noqa: BLE001
        return ""


def _line_metrics(text: str) -> tuple[int, int]:
    """(非空行数, 唯一 ``[n]`` 文献标记数)。"""
    lines = [ln for ln in text.splitlines() if ln.strip()]
    marks = {m for m in _REF_MARK_RE.findall(text)}
    return len(lines), len(marks)


def layout_report(
    pptx: str | Path,
    *,
    page: int,
    titles: list[int] | None = None,
    refs: list[tuple[int, int]] | None = None,
    tol: float = DEFAULT_TOL_IN,
) -> LayoutReport:
    """对一页做标题对齐与引用框规则体检（形状按 ID 指定）。"""
    from pptx import Presentation

    src = Path(pptx)
    if not src.exists():
        raise FileNotFoundError(f"pptx 不存在：{src}")
    prs = Presentation(str(src))
    if not 1 <= page <= len(prs.slides):
        raise ValueError(f"页码 {page} 越界（共 {len(prs.slides)} 页）")
    slide = prs.slides[page - 1]
    by_id = {int(sh.shape_id): sh for sh in slide.shapes}
    rep = LayoutReport(pptx=src, page=page)

    ids = list(titles or []) + [i for pair in (refs or []) for i in pair]
    missing = sorted({i for i in ids if i not in by_id})
    if missing:
        raise ValueError(
            "形状 ID 不存在于该页："
            + "、".join(str(i) for i in missing)
            + f"（现有 ID：{', '.join(str(i) for i in sorted(by_id))}）"
        )

    # —— ④ 焦点块标题对齐（上缘或左缘） ——
    if not titles:
        rep.skipped.append("标题对齐：未给 --title")
    elif len(titles) < 2:
        rep.skipped.append("标题对齐：需 ≥2 个标题形状")
    else:
        boxes = [(i, _box(by_id[i])) for i in titles]
        tops = [b[1] for _, b in boxes]
        lefts = [b[0] for _, b in boxes]
        d_top = max(tops) - min(tops)
        d_left = max(lefts) - min(lefts)
        ok = d_top <= tol or d_left <= tol
        parts = "；".join(
            f"{i} {by_id[i].name}(x={b[0]:.2f},top={b[1]:.2f})" for i, b in boxes
        )
        rep.checks.append(
            LayoutCheck(
                "焦点块标题对齐",
                ok,
                f"Δtop={d_top:.3f}in、Δleft={d_left:.3f}in（容差 {tol:.2f}；"
                f"上缘或左缘其一对齐即 PASS）：{parts}",
            )
        )

    # —— ③ 分焦点块引用框规则 ——
    if not refs:
        rep.skipped.append("引用框：未给 --refs REF:ANCHOR")
    else:
        tops: list[float] = []
        bottoms: list[float] = []
        for ref_id, anchor_id in refs:
            ref, anchor = by_id[ref_id], by_id[anchor_id]
            rb, ab = _box(ref), _box(anchor)
            delta = abs(rb[0] - ab[0])
            rep.checks.append(
                LayoutCheck(
                    f"引用框 {ref_id}↔{anchor_id} 左缘对齐",
                    delta <= tol,
                    f"Δleft={delta:.3f}in（容差 {tol:.2f}）",
                )
            )
            lines, marks = _line_metrics(_text(ref))
            if marks:
                ok = marks > 3 or lines == marks
                rule = "≥4 条不限行" if marks > 3 else "≤3 条一行一文献"
                rep.checks.append(
                    LayoutCheck(
                        f"引用框 {ref_id} 行数/文献数",
                        ok,
                        f"{marks} 条文献 / {lines} 行（{rule}）",
                    )
                )
            else:
                rep.skipped.append(f"引用框 {ref_id} 行数：文本无 [n] 标记")
            tops.append(rb[1])
            bottoms.append(rb[3])
        if len(refs) >= 2:
            d_top = max(tops) - min(tops)
            rep.checks.append(
                LayoutCheck(
                    "多引用框上缘一致",
                    d_top <= tol,
                    f"Δtop={d_top:.3f}in（容差 {tol:.2f}）",
                )
            )
            d_bottom = max(bottoms) - min(bottoms)
            rep.checks.append(
                LayoutCheck(
                    "多引用框下缘一致",
                    d_bottom <= tol,
                    f"Δbottom={d_bottom:.3f}in（容差 {tol:.2f}）",
                )
            )
    return rep
