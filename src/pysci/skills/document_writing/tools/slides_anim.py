"""``compose slides focus-anim``：按学院正式汇报预设生成焦点块动画 timing XML。

预设规格见 document-writing ``references/ppt-preset-formal-02-animation.md`` §5：

- 每块 = 1 个标签胶囊 + 若干正文面板/主图；**按形状 ID 逐个打靶**，不真组合；
- 聚焦（当前块）：胶囊填充与面板/主图描边 → ``#8496B0``，250 ms，ease-out；
- 失焦（非当前块）：胶囊填充 → ``#ADB9CA``、描边 → ``#D6DCE4``，1000 ms；
- 第 1 块随页自动（``onBegin tn=2``），其后每次单击前进一块（``delay=indefinite``）；
- 页面切换 fade 350 ms；
- 可选「进入·淡入」（presetID=10，250 ms），在块被聚焦时同刻播放。

本模块只做 XML 生成/注入，不渲染；生成后用 ``compose slides render/verify`` 做视觉终检。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from lxml import etree

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"

#: 预设 §5 默认值（学院正式汇报预设）。
DEFAULT_FOCUS_COLOR = "8496B0"
DEFAULT_BLUR_FILL = "ADB9CA"
DEFAULT_BLUR_LINE = "D6DCE4"
DEFAULT_FOCUS_MS = 250
DEFAULT_BLUR_MS = 1000
DEFAULT_TRANSITION_MS = 350
EASING_CHOICES = ("ease-out", "linear")

_HEX_RE = re.compile(r"^[0-9A-Fa-f]{6}$")


def _tag(ns: str, name: str) -> str:
    return f"{{{ns}}}{name}"


def _sub(parent: Any, ns: str, name: str, **attrs: str) -> Any:
    el = etree.SubElement(parent, _tag(ns, name))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def _hex(value: str) -> str:
    v = str(value or "").strip().lstrip("#").upper()
    if not _HEX_RE.match(v):
        raise ValueError(f"颜色须为 6 位十六进制（如 8496B0）：{value!r}")
    return v


@dataclass
class Block:
    """一个逻辑焦点块：胶囊（chip）+ 需要同步描边的形状（面板/主图）。"""

    chip: int
    shapes: list[int] = field(default_factory=list)
    fades: list[int] = field(default_factory=list)  # 聚焦本块时淡入的形状


@dataclass
class AnimOptions:
    """生成参数（默认值＝预设 §5）。"""

    focus_color: str = DEFAULT_FOCUS_COLOR
    blur_fill: str = DEFAULT_BLUR_FILL
    blur_line: str = DEFAULT_BLUR_LINE
    focus_ms: int = DEFAULT_FOCUS_MS
    blur_ms: int = DEFAULT_BLUR_MS
    transition_ms: int = DEFAULT_TRANSITION_MS
    easing: str = "ease-out"
    transition: bool = True


@dataclass
class AnimResult:
    """一次注入的读数。"""

    pptx: Path
    page: int
    blocks: int
    targets: int
    timing_groups: int
    transition: bool


def parse_block_specs(specs: list[str]) -> list[Block]:
    """解析 ``--block chip:shape1,shape2``（可重复；块序＝触发序）。

    ``chip`` 必填且为形状 ID；冒号后可省略（只播胶囊填充色）。ID 按出现序去重。
    """
    blocks: list[Block] = []
    seen_chips: set[int] = set()
    for raw in specs or []:
        spec = str(raw or "").strip()
        if not spec:
            raise ValueError("--block 不能为空；格式 chip:shape1,shape2")
        head, _, tail = spec.partition(":")
        try:
            chip = int(head.strip())
        except ValueError as exc:
            raise ValueError(f"--block 的胶囊 ID 非整数：{spec!r}") from exc
        if chip in seen_chips:
            raise ValueError(f"胶囊形状 ID {chip} 出现在多个块")
        seen_chips.add(chip)
        shapes: list[int] = []
        for piece in tail.split(","):
            piece = piece.strip()
            if not piece:
                continue
            try:
                sid = int(piece)
            except ValueError as exc:
                raise ValueError(f"--block 的形状 ID 非整数：{piece!r}") from exc
            if sid != chip and sid not in shapes:
                shapes.append(sid)
        blocks.append(Block(chip=chip, shapes=shapes))
    if not blocks:
        raise ValueError("至少需要一个 --block（块序＝触发序）")
    return blocks


def parse_fade_specs(specs: list[str], blocks: list[Block]) -> None:
    """解析 ``--fade 块序号:形状ID[,形状ID]``（块序号 1-based），写入 ``blocks``。"""
    for raw in specs or []:
        spec = str(raw or "").strip()
        head, _, tail = spec.partition(":")
        try:
            idx = int(head.strip())
        except ValueError as exc:
            raise ValueError(f"--fade 的块序号非整数：{spec!r}") from exc
        if not 1 <= idx <= len(blocks):
            raise ValueError(f"--fade 块序号 {idx} 越界（共 {len(blocks)} 块）")
        ids: list[int] = []
        for piece in tail.split(","):
            piece = piece.strip()
            if not piece:
                continue
            try:
                sid = int(piece)
            except ValueError as exc:
                raise ValueError(f"--fade 的形状 ID 非整数：{piece!r}") from exc
            if sid not in ids:
                ids.append(sid)
        if not ids:
            raise ValueError(f"--fade {idx}: 缺少形状 ID")
        blocks[idx - 1].fades.extend(i for i in ids if i not in blocks[idx - 1].fades)


class _Ids:
    """timing 内唯一 id 发生器（1/2 固定给 tmRoot/mainSeq）。"""

    def __init__(self) -> None:
        self.n = 3

    def next(self) -> int:
        value = self.n
        self.n += 1
        return value


def _set_attr(parent: Any, ids: _Ids, spid: int, attr: str, val: str, dur: int) -> None:
    """``p:set``：把形状的某个属性钉到定值（fill.on / stroke.on 等）。"""
    node = _sub(parent, P_NS, "set")
    bhvr = _sub(node, P_NS, "cBhvr")
    _sub(bhvr, P_NS, "cTn", id=ids.next(), dur=dur, fill="hold")
    tgt = _sub(bhvr, P_NS, "tgtEl")
    _sub(tgt, P_NS, "spTgt", spid=spid)
    names = _sub(bhvr, P_NS, "attrNameLst")
    _sub(names, P_NS, "attrName").text = attr
    _sub(node, P_NS, "to")
    to = node.find(_tag(P_NS, "to"))
    _sub(to, P_NS, "strVal", val=val)


def _color_effect(
    parent: Any,
    ids: _Ids,
    *,
    preset: int,
    attr: str,
    color: str,
    dur: int,
    spid: int,
    node_type: str,
    easing: str,
    fill: bool,
) -> None:
    """一条「强调·颜色」效果（presetID=1 填充 / 7 线条，emphesis sub=2）。"""
    par = _sub(parent, P_NS, "par")
    ctn = _sub(
        par,
        P_NS,
        "cTn",
        id=ids.next(),
        presetID=preset,
        presetClass="emph",
        presetSubtype=2,
        fill="hold",
        nodeType=node_type,
    )
    if easing == "ease-out":
        ctn.set("decel", "100000")
    conds = _sub(ctn, P_NS, "stCondLst")
    _sub(conds, P_NS, "cond", delay=0)
    children = _sub(ctn, P_NS, "childTnLst")
    anim = _sub(children, P_NS, "animClr", clrSpc="rgb", dir="cw")
    bhvr = _sub(anim, P_NS, "cBhvr")
    _sub(bhvr, P_NS, "cTn", id=ids.next(), dur=dur, fill="hold")
    tgt = _sub(bhvr, P_NS, "tgtEl")
    _sub(tgt, P_NS, "spTgt", spid=spid)
    names = _sub(bhvr, P_NS, "attrNameLst")
    _sub(names, P_NS, "attrName").text = attr
    to = _sub(anim, P_NS, "to")
    _sub(to, A_NS, "srgbClr", val=color)
    if fill:
        _set_attr(children, ids, spid, "fill.type", "solid", dur)
        _set_attr(children, ids, spid, "fill.on", "true", dur)
    else:
        _set_attr(children, ids, spid, "stroke.on", "true", dur)


def _entrance_effect(
    parent: Any, ids: _Ids, *, spid: int, dur: int, node_type: str
) -> None:
    """一条「进入·淡入」（presetID=10，entr sub=0）：聚焦同刻可见 + fade。"""
    par = _sub(parent, P_NS, "par")
    ctn = _sub(
        par,
        P_NS,
        "cTn",
        id=ids.next(),
        presetID=10,
        presetClass="entr",
        presetSubtype=0,
        fill="hold",
        nodeType=node_type,
    )
    conds = _sub(ctn, P_NS, "stCondLst")
    _sub(conds, P_NS, "cond", delay=0)
    children = _sub(ctn, P_NS, "childTnLst")
    _set_attr(children, ids, spid, "style.visibility", "visible", 1)
    anim = _sub(children, P_NS, "animEffect", transition="in", filter="fade")
    bhvr = _sub(anim, P_NS, "cBhvr")
    _sub(bhvr, P_NS, "cTn", id=ids.next(), dur=dur, fill="hold")
    tgt = _sub(bhvr, P_NS, "tgtEl")
    _sub(tgt, P_NS, "spTgt", spid=spid)


def _focus_effects(
    parent: Any,
    ids: _Ids,
    block: Block,
    *,
    opts: AnimOptions,
    first_node: str = "withEffect",
) -> None:
    """一块的聚焦效果：胶囊填充 + 各形状描边（均 withEffect，同时开始）。"""
    _color_effect(
        parent,
        ids,
        preset=1,
        attr="fillcolor",
        color=opts.focus_color,
        dur=opts.focus_ms,
        spid=block.chip,
        node_type=first_node,
        easing=opts.easing,
        fill=True,
    )
    for sid in block.shapes:
        _color_effect(
            parent,
            ids,
            preset=7,
            attr="stroke.color",
            color=opts.focus_color,
            dur=opts.focus_ms,
            spid=sid,
            node_type="withEffect",
            easing=opts.easing,
            fill=False,
        )
    for sid in block.fades:
        _entrance_effect(
            parent, ids, spid=sid, dur=opts.focus_ms, node_type="withEffect"
        )


def _blur_effects(
    parent: Any, ids: _Ids, block: Block, *, opts: AnimOptions, first_node: str
) -> None:
    """一块的失焦效果：胶囊填充、各形状描边（1000 ms，同时开始）。"""
    _color_effect(
        parent,
        ids,
        preset=1,
        attr="fillcolor",
        color=opts.blur_fill,
        dur=opts.blur_ms,
        spid=block.chip,
        node_type=first_node,
        easing=opts.easing,
        fill=True,
    )
    for sid in block.shapes:
        _color_effect(
            parent,
            ids,
            preset=7,
            attr="stroke.color",
            color=opts.blur_line,
            dur=opts.blur_ms,
            spid=sid,
            node_type="withEffect",
            easing=opts.easing,
            fill=False,
        )


def _remove(slide_el: Any, tag: str) -> None:
    for el in slide_el.findall(tag):
        slide_el.remove(el)


def _clear_timing(slide_el: Any) -> None:
    _remove(slide_el, _tag(P_NS, "timing"))


def _clear_transition(slide_el: Any) -> None:
    _remove(slide_el, _tag(P_NS, "transition"))
    for el in list(slide_el.findall(_tag(MC_NS, "AlternateContent"))):
        if el.find(f".//{_tag(P_NS, 'transition')}") is not None:
            slide_el.remove(el)


def _append_in_schema_order(slide_el: Any, el: Any) -> None:
    """``p:transition``/``p:timing`` 须在 ``p:extLst`` 之前。"""
    ext = slide_el.find(_tag(P_NS, "extLst"))
    if ext is not None:
        ext.addprevious(el)
    else:
        slide_el.append(el)


def _build_transition(dur: int) -> Any:
    alt = etree.Element(
        _tag(MC_NS, "AlternateContent"), nsmap={"mc": MC_NS, "p14": P14_NS}
    )
    choice = _sub(alt, MC_NS, "Choice", Requires="p14")
    tr = _sub(choice, P_NS, "transition")
    tr.set(_tag(P14_NS, "dur"), str(dur))
    _sub(tr, P_NS, "fade")
    fallback = _sub(alt, MC_NS, "Fallback")
    tr2 = _sub(fallback, P_NS, "transition")
    _sub(tr2, P_NS, "fade")
    return alt


def build_timing(slide_el: Any, blocks: list[Block], opts: AnimOptions) -> int:
    """把 timing 树写入 slide 元素（先清旧 timing），返回顶层组数。"""
    _clear_timing(slide_el)
    if opts.transition:
        _clear_transition(slide_el)
        _append_in_schema_order(slide_el, _build_transition(opts.transition_ms))
    timing = _sub(slide_el, P_NS, "timing")
    slide_el.remove(timing)  # 统一经 schema 序插回（transition 之前/之后固定）
    _append_in_schema_order(slide_el, timing)
    ids = _Ids()
    tn = _sub(timing, P_NS, "tnLst")
    root = _sub(tn, P_NS, "par")
    root_ctn = _sub(
        root, P_NS, "cTn", id=1, dur="indefinite", restart="never", nodeType="tmRoot"
    )
    root_children = _sub(root_ctn, P_NS, "childTnLst")
    seq = _sub(root_children, P_NS, "seq", concurrent=1, nextAc="seek")
    seq_ctn = _sub(seq, P_NS, "cTn", id=2, dur="indefinite", nodeType="mainSeq")
    groups = _sub(seq_ctn, P_NS, "childTnLst")

    # 第 1 块：随页自动（onBegin），所有效果 withEffect。
    first = _sub(groups, P_NS, "par")
    first_ctn = _sub(first, P_NS, "cTn", id=ids.next(), fill="hold")
    conds = _sub(first_ctn, P_NS, "stCondLst")
    _sub(conds, P_NS, "cond", delay="indefinite")
    begin = _sub(conds, P_NS, "cond", evt="onBegin", delay=0)
    _sub(begin, P_NS, "tn", val=2)
    wrap = _sub(first_ctn, P_NS, "childTnLst")
    inner = _sub(wrap, P_NS, "par")
    inner_ctn = _sub(inner, P_NS, "cTn", id=ids.next(), fill="hold")
    inner_conds = _sub(inner_ctn, P_NS, "stCondLst")
    _sub(inner_conds, P_NS, "cond", delay=0)
    effects = _sub(inner_ctn, P_NS, "childTnLst")
    _focus_effects(effects, ids, blocks[0], opts=opts)

    # 其后每块：一次单击 = 旧块失焦 + 新块聚焦（旧块填充是 clickEffect）。
    for i in range(1, len(blocks)):
        grp = _sub(groups, P_NS, "par")
        grp_ctn = _sub(grp, P_NS, "cTn", id=ids.next(), fill="hold")
        gconds = _sub(grp_ctn, P_NS, "stCondLst")
        _sub(gconds, P_NS, "cond", delay="indefinite")
        gwrap = _sub(grp_ctn, P_NS, "childTnLst")
        gi = _sub(gwrap, P_NS, "par")
        gi_ctn = _sub(gi, P_NS, "cTn", id=ids.next(), fill="hold")
        gconds2 = _sub(gi_ctn, P_NS, "stCondLst")
        _sub(gconds2, P_NS, "cond", delay=0)
        effs = _sub(gi_ctn, P_NS, "childTnLst")
        _blur_effects(effs, ids, blocks[i - 1], opts=opts, first_node="clickEffect")
        _focus_effects(effs, ids, blocks[i], opts=opts)
    return len(blocks)


def _validate_shape_ids(slide: Any, blocks: list[Block]) -> int:
    live = {int(sh.shape_id) for sh in slide.shapes}
    wanted: list[int] = []
    for b in blocks:
        wanted.extend([b.chip, *b.shapes, *b.fades])
    missing = sorted({sid for sid in wanted if sid not in live})
    if missing:
        raise ValueError(
            "形状 ID 不存在于该页："
            + "、".join(str(i) for i in missing)
            + f"（本页现有 ID：{', '.join(str(i) for i in sorted(live))}）"
        )
    return len(set(wanted))


def apply_focus_anim(
    pptx: str | Path,
    *,
    page: int,
    blocks: list[Block],
    opts: AnimOptions | None = None,
    out: str | Path | None = None,
) -> AnimResult:
    """把焦点块 timing 注入 ``page``（1-based），返回读数。

    Args:
        pptx: 源 .pptx。
        page: 页码（1-based）。
        blocks: 块序＝触发序（首块 onBegin，其余单击）。
        opts: 颜色/时长/缓动/页面切换参数；缺省取预设 §5。
        out: 输出路径；缺省原文件就地更新。

    Raises:
        FileNotFoundError: 文件不存在。
        ValueError: 页码越界 / 形状 ID 缺失。
    """
    from pptx import Presentation

    src = Path(pptx)
    if not src.exists():
        raise FileNotFoundError(f"pptx 不存在：{src}")
    opts = opts or AnimOptions()
    opts.focus_color = _hex(opts.focus_color)
    opts.blur_fill = _hex(opts.blur_fill)
    opts.blur_line = _hex(opts.blur_line)
    if opts.easing not in EASING_CHOICES:
        raise ValueError(f"缓动须为 {'/'.join(EASING_CHOICES)}：{opts.easing!r}")
    for name in ("focus_ms", "blur_ms", "transition_ms"):
        if int(getattr(opts, name)) <= 0:
            raise ValueError(f"{name} 须为正整数：{getattr(opts, name)!r}")
    prs = Presentation(str(src))
    if not 1 <= page <= len(prs.slides):
        raise ValueError(f"页码 {page} 越界（共 {len(prs.slides)} 页）")
    slide = prs.slides[page - 1]
    targets = _validate_shape_ids(slide, blocks)
    groups = build_timing(slide.element, blocks, opts)
    dest = Path(out) if out else src
    prs.save(str(dest))
    return AnimResult(
        pptx=dest,
        page=page,
        blocks=len(blocks),
        targets=targets,
        timing_groups=groups,
        transition=opts.transition,
    )
