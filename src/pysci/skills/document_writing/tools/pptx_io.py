"""pptx_io —— 用 python-pptx 结构化读取 .pptx。

相比 markitdown 的「拍平」提取，本模块按**幻灯片**组织，保留：
- 每页标题、正文（按段落/项目符号层级）
- 表格（渲染为 Markdown 表格）
- 图片清单（名称/替代文本/尺寸；可选导出为文件供 LLM「看图」）
- **演讲者备注**（旧汇报的关键信息常藏在这里）

这是「读懂历史汇报 pptx」的主力工具。仅支持 .pptx（OOXML），不支持旧版 .ppt。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

# python-pptx 是 [writing] 可选依赖；缺失时给出清晰指引而非裸 ImportError。
try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches
except ImportError as e:  # pragma: no cover
    raise ImportError(
        "python-pptx 未安装。请运行：uv sync --extra writing"
    ) from e


# ---------------------------------------------------------------------------
# 数据结构
# ---------------------------------------------------------------------------
@dataclass
class SlideContent:
    """单页幻灯片的结构化内容。"""

    index: int  # 从 1 开始
    layout: str = ""
    title: str = ""
    paragraphs: list[str] = field(default_factory=list)  # 正文段落（保留项目符号前缀）
    tables: list[list[list[str]]] = field(default_factory=list)  # 每个表：行→列→单元格
    images: list[dict[str, Any]] = field(default_factory=list)  # {name, alt, w, h, path?}
    notes: str = ""

    @property
    def text_char_count(self) -> int:
        return len(self.title) + sum(len(p) for p in self.paragraphs) + len(self.notes)


# ---------------------------------------------------------------------------
# 形状遍历（递归进 group）
# ---------------------------------------------------------------------------
def _iter_shapes(shapes: Any) -> list[Any]:
    """展开形状树：GroupShape 递归其子形状，其余原样返回。"""
    out: list[Any] = []
    for sh in shapes:
        if getattr(sh, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            try:
                out.extend(_iter_shapes(sh.shapes))
                continue
            except Exception:
                pass
        out.append(sh)
    return out


def _text_frame_paragraphs(tf: Any) -> list[str]:
    """把 text_frame 的段落抽成带项目符号层级的字符串列表。"""
    lines: list[str] = []
    for para in tf.paragraphs:
        text = "".join(run.text for run in para.runs).strip()
        if not text:
            # 有些段落 run 为空但 text 有值
            text = (para.text or "").strip()
        if not text:
            continue
        level = getattr(para, "level", 0) or 0
        prefix = "  " * level + ("• " if level or len(lines) else "")
        lines.append(f"{prefix}{text}")
    return lines


def _table_to_rows(table: Any) -> list[list[str]]:
    rows: list[list[str]] = []
    for r in table.rows:
        rows.append([(c.text or "").strip().replace("\n", " ") for c in r.cells])
    return rows


def _extract_image(shape: Any, idx: int) -> dict[str, Any]:
    info: dict[str, Any] = {"name": getattr(shape, "name", "") or f"image_{idx}"}
    try:
        img = shape.image
        info["content_type"] = img.content_type
        info["filename"] = img.filename
        info["size_bytes"] = len(img.blob)
    except Exception:
        pass
    try:
        info["w_emu"] = shape.width
        info["h_emu"] = shape.height
    except Exception:
        pass
    # 替代文本（alt text）常含语义信息
    try:
        alt = shape._element._nvXxPr.cNvPr.get("descr")  # noqa: SLF001
        if alt:
            info["alt"] = alt
    except Exception:
        pass
    return info


def _extract_picture_blob(shape: Any) -> tuple[str, bytes] | None:
    """取出图片的 (扩展名, 二进制)，供导出。"""
    try:
        img = shape.image
        ext = img.ext or "png"
        return ext, img.blob
    except Exception:
        return None


# ---------------------------------------------------------------------------
# 主入口：读取整个演示文稿
# ---------------------------------------------------------------------------
def read_pptx(path: str | Path, *, export_images_to: str | Path | None = None) -> list[SlideContent]:
    """读取 .pptx，返回每页的 SlideContent 列表。

    Args:
        path: .pptx 文件路径。
        export_images_to: 若给定目录，则把每页图片导出为文件，并写入 SlideContent.images[i]['path']
            （便于 LLM 用 Read 工具「看图」理解历史汇报里的示意图）。
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"pptx 不存在：{path}")
    if path.suffix.lower() not in (".pptx", ".pptm"):
        raise ValueError(f"仅支持 .pptx（不支持旧版 .ppt）：{path.name}")

    export_dir = Path(export_images_to) if export_images_to else None
    if export_dir:
        export_dir.mkdir(parents=True, exist_ok=True)

    prs = Presentation(str(path))
    slides: list[SlideContent] = []

    for i, slide in enumerate(prs.slides, 1):
        sc = SlideContent(index=i)
        try:
            sc.layout = slide.slide_layout.name or ""
        except Exception:
            pass

        # 标题
        try:
            if slide.shapes.title is not None:
                sc.title = (slide.shapes.title.text or "").strip()
        except Exception:
            pass

        img_counter = 0
        for shape in _iter_shapes(slide.shapes):
            try:
                # 表格
                if getattr(shape, "has_table", False):
                    rows = _table_to_rows(shape.table)
                    if rows:
                        sc.tables.append(rows)
                    continue
                # 图片
                if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                    img_counter += 1
                    info = _extract_image(shape, img_counter)
                    if export_dir:
                        blob = _extract_picture_blob(shape)
                        if blob:
                            ext, data = blob
                            fname = f"slide{i:02d}_img{img_counter}.{ext}"
                            fpath = export_dir / fname
                            fpath.write_bytes(data)
                            info["path"] = str(fpath)
                    sc.images.append(info)
                    continue
                # 文本
                if getattr(shape, "has_text_frame", False):
                    paras = _text_frame_paragraphs(shape.text_frame)
                    # 首个非空文本若尚无标题，用作标题
                    if not sc.title and paras:
                        sc.title = paras[0].lstrip("• ").strip()
                        sc.paragraphs.extend(paras[1:])
                    else:
                        sc.paragraphs.extend(paras)
            except Exception:
                # 单个异常形状不致命，跳过继续
                continue

        # 演讲者备注
        try:
            if slide.has_notes_slide:
                sc.notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        except Exception:
            pass

        slides.append(sc)

    return slides


# ---------------------------------------------------------------------------
# 渲染为 Markdown / dict
# ---------------------------------------------------------------------------
def _rows_to_md_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(norm[0]) + " |", "|" + "|".join(["---"] * width) + "|"]
    for r in norm[1:]:
        out.append("| " + " | ".join(cell.replace("|", "\\|") for cell in r) + " |")
    return "\n".join(out)


def slides_to_markdown(
    slides: list[SlideContent], *, source_name: str = "", include_notes: bool = True
) -> str:
    """把 SlideContent 列表渲染为便于 LLM 阅读的 Markdown。"""
    parts: list[str] = []
    if source_name:
        parts.append(f"# PPTX 提取：{source_name}\n")
    parts.append(f"> 共 {len(slides)} 页。按页组织，含正文/表格/图片/演讲者备注。\n")

    for sc in slides:
        head = f"## 第 {sc.index} 页"
        if sc.title:
            head += f"：{sc.title}"
        if sc.layout:
            head += f"  \n_（版式：{sc.layout}）_"
        parts.append(head + "\n")

        body = [p for p in sc.paragraphs if p.strip() and p.lstrip("• ").strip() != sc.title]
        if body:
            parts.append("\n".join(body) + "\n")

        for t_i, rows in enumerate(sc.tables, 1):
            parts.append(f"**表格 {t_i}：**\n\n" + _rows_to_md_table(rows) + "\n")

        if sc.images:
            inv = []
            for im in sc.images:
                bits = [im.get("name", "image")]
                if im.get("alt"):
                    bits.append(f'alt="{im["alt"]}"')
                if im.get("content_type"):
                    bits.append(im["content_type"])
                if im.get("path"):
                    bits.append(f'→ {Path(im["path"]).name}')
                inv.append("  - " + " | ".join(bits))
            parts.append("**图片：**\n" + "\n".join(inv) + "\n")

        if include_notes and sc.notes:
            parts.append(f"> **演讲者备注：** {sc.notes}\n")

    return "\n".join(parts)


def pptx_to_markdown(
    path: str | Path,
    *,
    include_notes: bool = True,
    export_images_to: str | Path | None = None,
) -> str:
    """一步到位：读取 .pptx → Markdown。"""
    path = Path(path)
    slides = read_pptx(path, export_images_to=export_images_to)
    return slides_to_markdown(
        slides, source_name=path.name, include_notes=include_notes
    )


# ---------------------------------------------------------------------------
# 写入：从规格 / Markdown 生成 .pptx
# ---------------------------------------------------------------------------
@dataclass
class SlideSpec:
    """单页幻灯片的写入规格（SlideContent 的“写”方向）。"""

    title: str = ""
    bullets: list[str] = field(default_factory=list)  # 用 2 空格缩进表示层级
    notes: str = ""
    table: list[list[str]] | None = None  # 行→列；给定时用「仅标题」版式


def _bullet_level(line: str) -> tuple[int, str]:
    """把带缩进/项目符号前缀的行解析为 (层级, 纯文本)。"""
    stripped = line.lstrip(" ")
    level = (len(line) - len(stripped)) // 2
    text = stripped.lstrip("-*• ").strip()
    return min(level, 4), text


def _fill_bullets(tf: Any, bullets: list[str]) -> None:
    first = True
    for b in bullets:
        level, text = _bullet_level(b)
        if not text:
            continue
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.text = text
        para.level = level


def add_title_slide(prs: Any, title: str, subtitle: str = "") -> Any:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    if subtitle and len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle
    return slide


def add_content_slide(prs: Any, spec: SlideSpec) -> Any:
    """按规格追加一页内容幻灯片（正文/表格/备注）。"""
    layout = prs.slide_layouts[5 if spec.table else 1]
    slide = prs.slides.add_slide(layout)
    if spec.title:
        slide.shapes.title.text = spec.title
    if spec.bullets:
        if spec.table:
            box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.4), Inches(1.2))
            _fill_bullets(box.text_frame, spec.bullets)
        else:
            _fill_bullets(slide.placeholders[1].text_frame, spec.bullets)
    if spec.table:
        rows = spec.table
        n_r = len(rows)
        n_c = max(len(r) for r in rows)
        top = Inches(2.8 if spec.bullets else 1.6)
        gf = slide.shapes.add_table(n_r, n_c, Inches(0.8), top, Inches(8.4), Inches(0.4 * n_r))
        tbl = gf.table
        for r in range(n_r):
            for c in range(n_c):
                tbl.cell(r, c).text = rows[r][c] if c < len(rows[r]) else ""
    if spec.notes:
        slide.notes_slide.notes_text_frame.text = spec.notes
    return slide


def build_pptx(
    specs: list[SlideSpec],
    out_path: str | Path,
    *,
    deck_title: str | None = None,
    deck_subtitle: str = "",
) -> Path:
    """从 SlideSpec 列表新建 .pptx；deck_title 给定时先加一页标题页。"""
    prs = Presentation()
    if deck_title:
        add_title_slide(prs, deck_title, deck_subtitle)
    for sp in specs:
        add_content_slide(prs, sp)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def add_slide_to(path: str | Path, spec: SlideSpec) -> Path:
    """向已有 .pptx 追加一页并保存。"""
    path = Path(path)
    prs = Presentation(str(path))
    add_content_slide(prs, spec)
    prs.save(str(path))
    return path


def markdown_to_pptx(md_text: str, out_path: str | Path) -> Path:
    """把约定式 Markdown 大纲转成 .pptx（与 slides_to_markdown 构成读写闭环）。

    约定：`# ` = 标题页（首个）/分节页；`## ` = 内容页标题；`- `/`* `/`• `（可缩进）=
    项目符号；`> ` = 演讲者备注；管道表 = 表格。
    """
    md_text = md_text.lstrip("\ufeff")  # 剥离 BOM（Windows 工具常写入）
    specs: list[SlideSpec] = []
    deck_title: str | None = None
    deck_subtitle = ""
    cur: SlideSpec | None = None
    table_buf: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_buf
        if table_buf and cur is not None:
            cur.table = table_buf
        table_buf = []

    for raw in md_text.splitlines():
        s = raw.strip()
        if not s:
            continue
        if s.startswith("|") and s.endswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if all(c and set(c) <= set("-: ") for c in cells):
                continue  # 表分隔行
            table_buf.append(cells)
            continue
        flush_table()
        if s.startswith("# "):
            if deck_title is None:
                deck_title = s[2:].strip()
            else:
                cur = SlideSpec(title=s[2:].strip())
                specs.append(cur)
            continue
        if s.startswith("## "):
            cur = SlideSpec(title=s[3:].strip())
            specs.append(cur)
            continue
        if s.startswith(">"):
            note = s.lstrip("> ").strip()
            note = note.removeprefix("**演讲者备注：**").strip()
            if cur is not None:
                cur.notes = f"{cur.notes} {note}".strip() if cur.notes else note
            elif deck_title is not None:
                deck_subtitle = f"{deck_subtitle} {note}".strip() if deck_subtitle else note
            continue
        if cur is None:
            cur = SlideSpec()
            specs.append(cur)
        cur.bullets.append(raw.rstrip())
    flush_table()
    return build_pptx(specs, out_path, deck_title=deck_title, deck_subtitle=deck_subtitle)


# ---------------------------------------------------------------------------
# CLI 自测
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="pptx 结构化提取 CLI")
    parser.add_argument("pptx", help=".pptx 文件路径")
    parser.add_argument("--no-notes", action="store_true", help="不含演讲者备注")
    parser.add_argument("--export-images", default=None, help="把图片导出到此目录")
    args = parser.parse_args()

    md = pptx_to_markdown(
        args.pptx,
        include_notes=not args.no_notes,
        export_images_to=args.export_images,
    )
    print(md)
