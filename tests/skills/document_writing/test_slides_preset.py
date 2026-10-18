"""组会 PPT 预设方案的三项机械检查（tb-20261017-212027-4-1016ppt-2 A.a5）。

三检查并入 `compose slides geometry`（预设单源：document-writing ``references/ppt-preset.md``）：

- (a) 内容页标题文本须为「·…·」；首页豁免、无标题的整页复制页自然跳过；
- (b) 标题框宽须随文本适应（程序化估算 vs 实际框宽，±10%）；
- (c) 非整页图须有描边（``<a:ln>`` 缺失即报）；≥ 画布 95% 的整页图豁免。

正例 = 用户原件 `20260918.pptx` 的标题**零误报**（四页标题都被识别，断言非空）；三负例
各由临时构造的迷你 deck 制造。真实 deck 用例在 data/ 缺席时跳过（worktree 内无数据）。
"""

from __future__ import annotations

import json
import subprocess
from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx", reason="需要 python-pptx")
Image = pytest.importorskip("PIL.Image", reason="需要 Pillow")

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from pysci.skills.document_writing.tools import compose, slides_visual  # noqa: E402

_REL_0918 = "data/research/1_gain_ep/汇报PPT/20260918/20260918.pptx"
_REL_1016 = "data/research/1_gain_ep/汇报PPT/20261016/20261016.pptx"


def _project_data_root() -> Path:
    """含 data/ 的项目根；worktree 里 data/（gitignored）缺席时经 git 公共目录回退主树。"""
    from pysci.paths import PROJECT_ROOT

    if (PROJECT_ROOT / "data").is_dir():
        return PROJECT_ROOT
    try:
        out = subprocess.run(
            ["git", "rev-parse", "--path-format=absolute", "--git-common-dir"],
            cwd=PROJECT_ROOT,
            capture_output=True,
            text=True,
            check=True,
        ).stdout.strip()
        root = Path(out).parent
        if (root / "data").is_dir():
            return root
    except (OSError, subprocess.CalledProcessError):
        pass
    return PROJECT_ROOT


def _real_deck(rel: str) -> Path:
    p = _project_data_root() / rel
    if not p.exists():
        pytest.skip(f"缺少真实 deck（data/ 不在本树）：{rel}")
    return p


def _png(path: Path) -> Path:
    Image.new("RGB", (60, 40), (30, 90, 200)).save(path)
    return path


def _blank(prs: Presentation) -> None:
    prs.slides.add_slide(prs.slide_layouts[6])


def _titled(
    prs: Presentation,
    text: str,
    *,
    width: float | None = None,
    size: int = 20,
) -> tuple[object, float]:
    """在预设标题区 (0.60,0.44) 加标题文本框；width=None 时按估算宽自适应。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(
        Inches(0.60), Inches(0.44), Inches(6.01), Inches(0.44)
    )
    box.text_frame.text = text
    for run in box.text_frame.paragraphs[0].runs:
        run.font.size = Pt(size)
    est = slides_visual.estimate_title_width_in(box)
    box.width = Inches(width if width is not None else est)
    return slide, est


def _pic(
    slide,
    img: Path,
    *,
    left: float = 1.0,
    top: float = 1.5,
    width: float = 2.0,
    height: float = 1.5,
    border: bool = False,
) -> None:
    pic = slide.shapes.add_picture(
        str(img), Inches(left), Inches(top), Inches(width), Inches(height)
    )
    if border:
        pic.line.color.rgb = RGBColor(0, 0, 0)


def _save(prs: Presentation, tmp_path: Path, name: str = "deck.pptx") -> Path:
    out = tmp_path / name
    prs.save(str(out))
    return out


# ---------------------------------------------------------------------------
# 正例：用户原件标题零误报
# ---------------------------------------------------------------------------
def test_user_original_titles_have_zero_false_positives():
    """`20260918.pptx` 四页标题全部被识别，且形态/框宽零误报（非空断言防「没检测」）。"""
    deck = _real_deck(_REL_0918)
    prs = Presentation(str(deck))
    titles = {
        i: slides_visual._title_text(slides_visual._preset_title_shape(s))
        for i, s in enumerate(prs.slides, 1)
    }
    titles = {i: t for i, t in titles.items() if t}
    assert titles == {
        2: "·现状总结·",
        3: "·今天的AI能做什么·",
        4: "·一些准备·",
        5: "·pySci·",
    }, titles

    rep = slides_visual.geometry_report(deck)
    assert all(not s.title_form_issues for s in rep.slides), [
        s.title_form_issues for s in rep.slides if s.title_form_issues
    ]
    assert all(not s.title_width_issues for s in rep.slides), [
        s.title_width_issues for s in rep.slides if s.title_width_issues
    ]


def test_1016_deck_reports_expected_preset_violations():
    """修复前的 `20261016.pptx`：页 2/6 无「·」、页 2/3/6 框宽失配、页 2/6 图无描边。"""
    deck = _real_deck(_REL_1016)
    rep = slides_visual.geometry_report(deck)
    form = {s.index for s in rep.slides if s.title_form_issues}
    width = {s.index for s in rep.slides if s.title_width_issues}
    border = {s.index for s in rep.slides if s.picture_border_issues}
    assert form == {2, 6}, form
    assert width == {2, 3, 6}, width
    assert border == {2, 6}, border  # 页 4/5/7 整页满幅图豁免


# ---------------------------------------------------------------------------
# 三负例（合成 deck）+ 正例对照
# ---------------------------------------------------------------------------
def test_clean_preset_deck_has_no_issues(tmp_path):
    """标题「·…·」且框宽自适应、图带描边 → 三项零报（首页豁免亦被覆盖）。"""
    prs = Presentation()
    _blank(prs)  # 页 1（首页）
    slide, est = _titled(prs, "·现状总结·")
    _pic(slide, _png(tmp_path / "a.png"), border=True)
    deck = _save(prs, tmp_path)

    rep = slides_visual.geometry_report(deck)
    assert est == pytest.approx(1.75, abs=0.02), est  # 对标 20260918 的 1.78 in
    for s in rep.slides:
        assert not s.title_form_issues and not s.title_width_issues, s
        assert not s.picture_border_issues, s
    assert rep.ok, rep.summary()


def test_missing_dot_form_is_flagged(tmp_path):
    """负例 (a)：内容页标题无前后「·」→ 形态违规。"""
    prs = Presentation()
    _blank(prs)
    _titled(prs, "无装饰标题")
    deck = _save(prs, tmp_path)
    sg = slides_visual.geometry_report(deck).slides[1]
    assert sg.title_form_issues and "·…·" in sg.title_form_issues[0]
    assert not sg.title_width_issues, sg.title_width_issues
    assert sg.violations >= 1


def test_overwide_title_box_is_flagged(tmp_path):
    """负例 (b)：短标题被统一拉成 6.01 in 框宽 → 适配违规（预设的返工事故原型）。"""
    prs = Presentation()
    _blank(prs)
    _titled(prs, "·pySci·", width=6.01)
    deck = _save(prs, tmp_path)
    sg = slides_visual.geometry_report(deck).slides[1]
    assert sg.title_width_issues and "6.01 in" in sg.title_width_issues[0]
    assert not sg.title_form_issues, sg.title_form_issues


def test_picture_without_border_is_flagged(tmp_path):
    """负例 (c)：非整页图无 <a:ln> → 描边违规。"""
    prs = Presentation()
    _blank(prs)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _pic(slide, _png(tmp_path / "b.png"), border=False)
    deck = _save(prs, tmp_path)
    sg = slides_visual.geometry_report(deck).slides[1]
    assert sg.picture_border_issues and "无描边" in sg.picture_border_issues[0]


def test_full_page_picture_is_border_exempt(tmp_path):
    """整页满幅图（= 画布尺寸）免描边检查。"""
    prs = Presentation()
    _blank(prs)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _pic(
        slide,
        _png(tmp_path / "bg.png"),
        left=0.0,
        top=0.0,
        width=10.0,
        height=7.5,
        border=False,
    )
    deck = _save(prs, tmp_path)
    assert slides_visual.geometry_report(deck).slides[1].picture_border_issues == []


def test_first_page_and_explicit_exemption(tmp_path):
    """首页默认豁免；`title_exempt_pages` 可把指定页一并豁免。"""
    prs = Presentation()
    _titled(prs, "首页标题允许无点")  # 第 1 页：首页豁免
    _titled(prs, "第二页无点")
    deck = _save(prs, tmp_path)

    rep = slides_visual.geometry_report(deck)
    assert rep.slides[0].title_form_issues == []
    assert rep.slides[1].title_form_issues

    rep2 = slides_visual.geometry_report(deck, title_exempt_pages={1, 2})
    assert all(not s.title_form_issues for s in rep2.slides)


# ---------------------------------------------------------------------------
# CLI / JSON / 退出码
# ---------------------------------------------------------------------------
def test_cli_strict_fails_on_preset_violation(tmp_path):
    prs = Presentation()
    _blank(prs)
    _titled(prs, "无装饰标题")
    deck = _save(prs, tmp_path)
    assert compose.main(["slides", "geometry", str(deck)]) == 0  # 默认只报告
    assert compose.main(["slides", "geometry", str(deck), "--strict"]) == 1


def test_cli_json_exposes_preset_issues(tmp_path, capsys):
    prs = Presentation()
    _blank(prs)
    _titled(prs, "无装饰标题")
    deck = _save(prs, tmp_path)
    assert compose.main(["slides", "geometry", str(deck), "--json"]) == 0
    data = json.loads(capsys.readouterr().out)
    assert data["slides"][1]["title_form_issues"]
    assert "title_width_issues" in data["slides"][1]
    assert "picture_border_issues" in data["slides"][1]
