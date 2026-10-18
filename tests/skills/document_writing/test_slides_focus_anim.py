"""`compose slides focus-anim`：按学院预设生成焦点块 timing（tb-5 证据条目②）。

判据（预设分册 02 §5）：首块 onBegin 自动；其后单击逐块；聚焦 250 ms / 失焦 1000 ms；
decel=100000；页面 fade 350 ms；只打靶形状 ID，不真组合。这里解包 XML 逐项钉住。
"""

from __future__ import annotations

from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx", reason="需要 python-pptx")

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches  # noqa: E402

from pysci.skills.document_writing.tools import compose  # noqa: E402

P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"


def _deck(tmp_path: Path) -> tuple[Path, dict[str, int]]:
    """三块形状的迷你 deck：每块一个胶囊 + 面板，另有一张待淡入图。"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(0.6))
    tb.text_frame.text = "焦点块动画测试页"

    def rect(x: float):
        return slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(2), Inches(1), Inches(1)
        )

    chip1, p1, p2, chip2, p3, chip3, p4, fade = (
        rect(0.5),
        rect(2.0),
        rect(3.0),
        rect(4.0),
        rect(5.0),
        rect(6.0),
        rect(7.0),
        rect(8.0),
    )
    deck = tmp_path / "focus.pptx"
    prs.save(str(deck))
    return deck, {
        "chip1": chip1.shape_id,
        "p1": p1.shape_id,
        "p2": p2.shape_id,
        "chip2": chip2.shape_id,
        "p3": p3.shape_id,
        "chip3": chip3.shape_id,
        "p4": p4.shape_id,
        "fade": fade.shape_id,
    }


def _slide(path: Path):
    return Presentation(str(path)).slides[0]


def _groups(slide) -> list:
    timing = slide.element.find(qn("p:timing"))
    assert timing is not None, "未写入 <p:timing>"
    return timing.findall(
        f".//{qn('p:cTn')}[@nodeType='mainSeq']/{qn('p:childTnLst')}/{qn('p:par')}"
    )


def _colors(group) -> list[tuple[str, str, int, int, str, str | None]]:
    """组内全部颜色效果：(attr, color, spid, dur, nodeType, decel)。"""
    out = []
    for anim in group.findall(f".//{qn('p:animClr')}"):
        bhvr = anim.find(qn("p:cBhvr"))
        attr = bhvr.find(f"{qn('p:attrNameLst')}/{qn('p:attrName')}").text
        color = anim.find(f"{qn('p:to')}/{qn('a:srgbClr')}").get("val")
        spid = int(bhvr.find(f"{qn('p:tgtEl')}/{qn('p:spTgt')}").get("spid"))
        ctn = bhvr.find(qn("p:cTn"))
        parent = next(anim.iterancestors(qn("p:cTn")))
        out.append(
            (
                attr,
                color,
                spid,
                int(ctn.get("dur")),
                parent.get("nodeType"),
                parent.get("decel"),
            )
        )
    return out


def _run(deck: Path, *extra: str) -> int:
    s = _deck_ids(deck)
    return compose.main(
        [
            "slides",
            "focus-anim",
            str(deck),
            "--page",
            "1",
            "--block",
            f"{s['chip1']}:{s['p1']},{s['p2']}",
            "--block",
            f"{s['chip2']}:{s['p3']}",
            "--block",
            f"{s['chip3']}:{s['p4']}",
            *extra,
        ]
    )


def _deck_ids(deck: Path) -> dict[str, int]:
    prs = Presentation(str(deck))
    names = ["chip1", "p1", "p2", "chip2", "p3", "chip3", "p4", "fade"]
    shapes = list(prs.slides[0].shapes)[1:]  # 第 1 个形状是标题文本框
    return {n: sh.shape_id for n, sh in zip(names, shapes)}


def test_generates_focus_and_blur_timing(tmp_path):
    deck, ids = _deck(tmp_path)
    res = compose.main(
        [
            "slides",
            "focus-anim",
            str(deck),
            "--page",
            "1",
            "--block",
            f"{ids['chip1']}:{ids['p1']},{ids['p2']}",
            "--block",
            f"{ids['chip2']}:{ids['p3']}",
            "--block",
            f"{ids['chip3']}:{ids['p4']}",
        ]
    )
    assert res == 0
    slide = _slide(deck)
    groups = _groups(slide)
    assert len(groups) == 3

    # 首块：onBegin 自动（tn=2），全部 withEffect
    first = groups[0].find(f"{qn('p:cTn')}/{qn('p:stCondLst')}")
    begin = first.find(f"{qn('p:cond')}[@evt='onBegin']/{qn('p:tn')}")
    assert begin is not None and begin.get("val") == "2"
    g0 = _colors(groups[0])
    assert {(a, c, s, d) for a, c, s, d, *_ in g0} == {
        ("fillcolor", "8496B0", ids["chip1"], 250),
        ("stroke.color", "8496B0", ids["p1"], 250),
        ("stroke.color", "8496B0", ids["p2"], 250),
    }
    assert all(node == "withEffect" and decel == "100000" for *_, node, decel in g0)

    # 其后单击：旧块失焦 1000 + 新块聚焦 250；首效果 clickEffect
    assert (
        groups[1].find(f"{qn('p:cTn')}/{qn('p:stCondLst')}/{qn('p:cond')}").get("delay")
        == "indefinite"
    )
    g1 = _colors(groups[1])
    assert ("fillcolor", "ADB9CA", ids["chip1"], 1000, "clickEffect", "100000") in g1
    assert ("stroke.color", "D6DCE4", ids["p1"], 1000, "withEffect", "100000") in g1
    assert ("fillcolor", "8496B0", ids["chip2"], 250, "withEffect", "100000") in g1


def test_transition_fade_350_and_xml_order(tmp_path):
    deck, ids = _deck(tmp_path)
    assert _run(deck) == 0
    slide = _slide(deck)
    trans = slide.element.find(f".//{qn('p:transition')}")
    assert trans is not None
    assert trans.get(f"{{{P14}}}dur") == "350"
    assert trans.find(qn("p:fade")) is not None
    tags = [c.tag for c in slide.element]
    assert tags.index(qn("p:timing")) > next(
        i for i, t in enumerate(tags) if t.endswith("}AlternateContent")
    )


def test_rerun_is_idempotent(tmp_path):
    deck, ids = _deck(tmp_path)
    assert _run(deck) == 0
    assert _run(deck) == 0
    slide = _slide(deck)
    assert len(slide.element.findall(qn("p:timing"))) == 1
    alts = [
        el
        for el in slide.element
        if el.tag.endswith("}AlternateContent")
        and el.find(f".//{qn('p:transition')}") is not None
    ]
    assert len(alts) == 1, "页面切换只应有一份 mc:AlternateContent"
    assert len(_groups(slide)) == 3


def test_custom_colors_durations_and_linear_easing(tmp_path):
    deck, ids = _deck(tmp_path)
    assert (
        _run(
            deck,
            "--focus-color",
            "112233",
            "--blur-fill",
            "445566",
            "--blur-line",
            "778899",
            "--focus-ms",
            "300",
            "--blur-ms",
            "800",
            "--easing",
            "linear",
            "--no-transition",
        )
        == 0
    )
    slide = _slide(deck)
    groups = _groups(slide)
    g1 = _colors(groups[1])
    assert ("fillcolor", "445566", ids["chip1"], 800, "clickEffect", None) in g1
    assert ("fillcolor", "112233", ids["chip2"], 300, "withEffect", None) in g1
    assert slide.element.find(f".//{qn('p:transition')}") is None
    assert slide.element.find(qn("p:timing")) is not None


def test_entrance_fade_targets_shape(tmp_path):
    deck, ids = _deck(tmp_path)
    assert _run(deck, "--fade", f"2:{ids['fade']}") == 0
    slide = _slide(deck)
    groups = _groups(slide)
    entrs = groups[1].findall(f".//{qn('p:cTn')}[@presetClass='entr']")
    assert len(entrs) == 1 and entrs[0].get("presetID") == "10"
    anim = entrs[0].find(f".//{qn('p:animEffect')}")
    assert anim.get("filter") == "fade" and anim.get("transition") == "in"
    assert anim.find(f".//{qn('p:spTgt')}").get("spid") == str(ids["fade"])
    assert entrs[0].find(f".//{qn('p:cTn')}[@dur='250']") is not None


def test_unknown_shape_id_is_rejected_without_writing(tmp_path, capsys):
    deck, ids = _deck(tmp_path)
    before = deck.read_bytes()
    rc = compose.main(
        [
            "slides",
            "focus-anim",
            str(deck),
            "--page",
            "1",
            "--block",
            "99999:1",
        ]
    )
    assert rc == 2
    assert "形状 ID 不存在" in capsys.readouterr().err
    assert deck.read_bytes() == before, "校验失败不得落盘"


def test_bad_spec_errors(tmp_path):
    deck, _ = _deck(tmp_path)
    with pytest.raises(SystemExit) as exc:  # --block 为必填：argparse 报用法
        compose.main(["slides", "focus-anim", str(deck), "--page", "1"])
    assert exc.value.code == 2
    assert (
        compose.main(
            ["slides", "focus-anim", str(deck), "--page", "1", "--block", "x:1"]
        )
        == 2
    )
    assert (
        compose.main(["slides", "focus-anim", str(deck), "--page", "9", "--block", "1"])
        == 2
    )


def test_verify_passes_after_generation(tmp_path):
    deck, ids = _deck(tmp_path)
    assert _run(deck) == 0
    render = tmp_path / "render"
    rc = compose.main(["slides", "render", str(deck), "--out", str(render)])
    if rc == 3:
        pytest.skip("未检测到 LibreOffice/wpscli，渲染链不可用")
    assert rc == 0
    assert (
        compose.main(["slides", "verify", str(deck), "--render-dir", str(render)]) == 0
    )


def test_no_transition_preserves_existing_page_transition(tmp_path):
    """`--no-transition` 只表示不写切换，不擅自删除 deck 已有的页面切换。"""
    deck, ids = _deck(tmp_path)
    assert _run(deck) == 0
    assert _run(deck, "--no-transition") == 0
    slide = _slide(deck)
    assert slide.element.find(f".//{qn('p:transition')}") is not None
