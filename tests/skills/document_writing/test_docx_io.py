"""docx_io 测试：结构化读取 / 构建 / 增量追加（无需外部工具）。

markdown→docx 写转换已改由 Pandoc 承担，测试见 test_pandoc_convert.py。
"""

from __future__ import annotations

import pytest

docx = pytest.importorskip("docx", reason="需要 python-docx")

from pysci.skills.document_writing.tools import docx_io  # noqa: E402


def test_build_and_read_roundtrip(tmp_path):
    out = tmp_path / "doc.docx"
    blocks = [
        docx_io.DocxBlock(kind="heading", text="Intro", level=1),
        docx_io.DocxBlock(kind="paragraph", text="Body sentence."),
        docx_io.DocxBlock(kind="bullet", text="point one"),
        docx_io.DocxBlock(kind="bullet", text="sub point", level=1),
        docx_io.DocxBlock(kind="table", rows=[["g", "Q"], ["0.5", "820"]]),
    ]
    docx_io.build_docx(blocks, out, title="Report")

    got = docx_io.read_docx(out)
    kinds = [b.kind for b in got]
    assert kinds[0] == "heading" and got[0].text == "Report"
    assert (
        "heading" in kinds
        and "paragraph" in kinds
        and "bullet" in kinds
        and "table" in kinds
    )
    tbl = next(b for b in got if b.kind == "table")
    assert tbl.rows[0][:2] == ["g", "Q"]
    sub = next(b for b in got if b.kind == "bullet" and b.text == "sub point")
    assert sub.level == 1


def test_add_block_appends(tmp_path):
    out = tmp_path / "doc.docx"
    docx_io.build_docx([docx_io.DocxBlock(kind="paragraph", text="first")], out)
    assert len(docx_io.read_docx(out)) == 1
    docx_io.add_block_to(
        out, docx_io.DocxBlock(kind="heading", text="Appended", level=2)
    )
    got = docx_io.read_docx(out)
    assert len(got) == 2
    assert got[1].kind == "heading" and got[1].text == "Appended"


# ---------------------------------------------------------------------------
# 内联 OMML 公式文本不得被静默丢弃（backlog 20261017-052625-writing）
# ---------------------------------------------------------------------------
from docx.oxml import parse_xml  # noqa: E402

_M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _omml(text: str):
    """造一个只含文本的内联公式元素 ``<m:oMath><m:r><m:t>…``。"""
    return parse_xml(
        f'<m:oMath xmlns:m="{_M_NS}"><m:r><m:t>{text}</m:t></m:r></m:oMath>'
    )


def _inject_math_paragraph(doc, before: str, math: str, after: str):
    p = doc.add_paragraph()
    if before:
        p.add_run(before)
    p._p.append(_omml(math))
    if after:
        p.add_run(after)
    return p


def test_inline_omml_formula_text_is_extracted(tmp_path):
    """正文内联公式（数值/单位常整段是公式）必须进入提取结果。"""
    doc = docx.Document()
    _inject_math_paragraph(
        doc, "该实验平台的主体由高度为 ", "4 cm", " 的二维上下平行刚性平板波导构成。"
    )
    out = tmp_path / "formula.docx"
    doc.save(str(out))

    para = next(b for b in docx_io.read_docx(out) if b.kind == "paragraph")
    assert "4 cm" in para.text, para.text
    assert "高度为 4 cm 的" in para.text, para.text
    # Markdown 出口同样保住公式文本
    assert "4 cm" in docx_io.docx_to_markdown(out)


def test_omml_inside_table_cell_is_extracted(tmp_path):
    """同一缺陷类在表格单元格里也要修（python-docx 的 ``cell.text`` 同样丢 OMML）。"""
    doc = docx.Document()
    tbl = doc.add_table(rows=1, cols=2)
    tbl.cell(0, 0).text = "频率"
    tbl.cell(0, 1).paragraphs[0]._p.append(_omml("3 kHz"))
    out = tmp_path / "cellmath.docx"
    doc.save(str(out))

    t = next(b for b in docx_io.read_docx(out) if b.kind == "table")
    assert t.rows[0][0] == "频率"
    assert t.rows[0][1] == "3 kHz", t.rows
