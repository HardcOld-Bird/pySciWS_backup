"""``compose docx marks``：run 级标注抽取（tb-20261017-211704-2-2 F2）。

机械对照面：writing 自检 6 / reviewer G10 要在「改稿标注」上核对作者声明的修改——
红删（strike + shd FF0000）与黄粗（shd FFFF00 + bold）必须能被逐 run 读出来。
负例：纯文本段不产出标注；缺一类标注（只有 strike 无底色）不得捏造缺失项。
"""

from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from pysci.skills.document_writing.tools import compose, docx_io


def _shade(run, fill: str) -> None:
    rpr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    rpr.append(shd)


def _sample(tmp_path):
    """样张：段 0 = 红删＋黄粗；段 1 = 纯文本；段 2 = 仅删除线（缺底色）。"""
    doc = Document()
    p0 = doc.add_paragraph()
    old = p0.add_run("旧句")
    old.font.strike = True
    _shade(old, "FF0000")
    new = p0.add_run("新句")
    new.font.bold = True
    _shade(new, "FFFF00")

    doc.add_paragraph("纯文本段：无任何标注")

    p2 = doc.add_paragraph()
    only = p2.add_run("仅删除线片段")
    only.font.strike = True

    path = tmp_path / "sample.docx"
    doc.save(str(path))
    return path


def test_extracts_strike_shd_and_bold(tmp_path):
    grouped, total = docx_io.extract_run_marks(_sample(tmp_path))
    assert total == 3, grouped
    by_para = {i: (text, runs) for i, text, runs in grouped}
    assert set(by_para) == {0, 2}, "纯文本段不得出现在结果里"
    runs0 = by_para[0][1]
    assert runs0[0].marks == ["strike", "shd=FF0000"]
    assert runs0[0].text == "旧句"
    assert runs0[1].marks == ["shd=FFFF00", "bold"]
    assert by_para[2][1][0].marks == ["strike"]


def test_cli_marks_output_is_greppable(tmp_path, capsys):
    rc = compose.main(["docx", "marks", str(_sample(tmp_path))])
    out = capsys.readouterr().out
    assert rc == 0
    for frag in (
        "docx marks:",
        "段 0",
        "strike shd=FF0000",
        "shd=FFFF00 bold",
        "汇总：",
    ):
        assert frag in out, out
    assert "旧句" in out and "新句" in out


def test_para_filter_narrows_to_matching_paragraph(tmp_path, capsys):
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("目标段文字")
    r.font.strike = True
    _shade(r, "FF0000")
    doc.add_paragraph("别段文字")
    path = tmp_path / "filter.docx"
    doc.save(str(path))

    rc = compose.main(["docx", "marks", str(path), "--para", "目标段"])
    out = capsys.readouterr().out
    assert rc == 0
    assert "目标段文字" in out and "别段文字" not in out, out
    assert "FF0000" in out and "strike" in out, out


def test_plain_text_docx_reports_no_marks(tmp_path, capsys):
    doc = Document()
    doc.add_paragraph("只有纯文本，没有任何标注")
    path = tmp_path / "plain.docx"
    doc.save(str(path))

    rc = compose.main(["docx", "marks", str(path)])
    out = capsys.readouterr().out
    assert rc == 0
    assert "（无标注 run）" in out, out
    assert "汇总：" not in out


def test_missing_annotation_is_not_fabricated(tmp_path, capsys):
    """负例：只有 strike 的样张，输出不得出现 shd/bold。"""
    path = _sample(tmp_path)
    rc = compose.main(["docx", "marks", str(path), "--para", "仅删除线"])
    out = capsys.readouterr().out
    assert rc == 0
    assert "strike" in out, out
    assert "shd=" not in out and "bold" not in out, out
