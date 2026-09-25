"""deck_digest 测试：用合成的小 deck 验证「去重导图 + 邻近文字 + 分节分块 + 骨架 + 断点账本」。

不依赖 LibreOffice（渲染为可选步骤，缺失时只在 progress 里记录说明）。
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from pysci.skills.document_writing.tools import deck_digest


# ---------------------------------------------------------------------------
# 合成 deck
# ---------------------------------------------------------------------------
def _png(color: tuple[int, int, int], size: tuple[int, int] = (32, 24)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    return buf.getvalue()


RED = _png((200, 30, 30))
BLUE = _png((30, 30, 200))


def _gif(colors: list[tuple[int, int, int]], size: tuple[int, int] = (32, 24)) -> bytes:
    """造一个多帧 gif（LLM 不能直读，用于验证抽帧预览）。"""
    frames = [Image.new("RGB", size, c) for c in colors]
    buf = io.BytesIO()
    frames[0].save(
        buf, format="GIF", save_all=True, append_images=frames[1:], duration=50, loop=0
    )
    return buf.getvalue()


@pytest.fixture()
def deck(tmp_path: Path) -> Path:
    """5 页：标题页 / 分节页 / 双图页（含图注+备注）/ 重复图页 / 纯文本页。"""
    prs = Presentation()
    lay = prs.slide_layouts

    s = prs.slides.add_slide(lay[0])  # Title Slide
    s.shapes.title.text = "Gain-EP 研究进展汇总"
    s.placeholders[1].text = "2018-2026"

    s = prs.slides.add_slide(lay[2])  # Section Header → 应被识别为分节页
    s.shapes.title.text = "第一部分 理论与仿真"

    s = prs.slides.add_slide(lay[5])  # Title Only + 两张图 + 图注 + 备注
    s.shapes.title.text = "COMSOL 声压场"
    s.shapes.add_picture(io.BytesIO(RED), Inches(0.5), Inches(1.6), Inches(4), Inches(3))
    s.shapes.add_picture(io.BytesIO(BLUE), Inches(5.0), Inches(1.6), Inches(4), Inches(3))
    cap = s.shapes.add_textbox(Inches(0.5), Inches(5.0), Inches(9), Inches(0.8))
    cap.text_frame.text = "Fig. 1 声压幅值分布（左：增益前；右：增益后）"
    s.notes_slide.notes_text_frame.text = "这页展示 EP 附近的场分布"

    s = prs.slides.add_slide(lay[1])  # 重复使用 RED → 去重
    s.shapes.title.text = "装置示意（重复图）"
    s.shapes.add_picture(io.BytesIO(RED), Inches(1.0), Inches(2.0), Inches(3), Inches(2))

    s = prs.slides.add_slide(lay[1])  # 纯文本页
    s.shapes.title.text = "小结"
    s.placeholders[1].text_frame.text = "增益-损耗平衡可调"

    out = tmp_path / "deck.pptx"
    prs.save(str(out))
    return out


# ---------------------------------------------------------------------------
# 用例
# ---------------------------------------------------------------------------
def test_digest_stats_and_dedup(deck: Path, tmp_path: Path):
    out = tmp_path / "translated"
    res = deck_digest.digest_pptx(deck, out)

    assert res.n_slides == 5
    assert res.n_pictures == 3
    assert res.n_unique_images == 2  # RED 重复一次
    assert res.n_duplicates == 1
    assert res.n_figure_slides == 2
    assert res.n_text_slides == 3
    assert len(list((out / "images").glob("*"))) == 2  # 去重后只落 2 个文件

    manifest = json.loads((out / "images_manifest.json").read_text(encoding="utf-8"))
    assert len(manifest) == 2
    dup = [m for m in manifest.values() if len(m["occurrences"]) > 1]
    assert len(dup) == 1
    assert dup[0]["occurrences"][0]["slide"] == 3
    assert dup[0]["occurrences"][1]["slide"] == 4


def test_sections_and_chunks(deck: Path, tmp_path: Path):
    out = tmp_path / "translated"
    res = deck_digest.digest_pptx(deck, out)

    titles = [s["title"] for s in res.sections]
    assert "第一部分 理论与仿真" in titles
    # 分节页在 slide 2 → 两节：1-1 与 2-5
    assert [(s["start"], s["end"]) for s in res.sections] == [(1, 1), (2, 5)]
    # 两节共 5 页 ≤ 默认单块上限 40 → 合并为 1 个 md 块
    assert len(res.chunks) == 1
    assert all(p.is_file() for p in res.chunks)


def test_small_max_chunk_splits_sections(deck: Path, tmp_path: Path):
    """单块上限很小时：超长的节内部再切，且块不跨节合并超限。"""
    out = tmp_path / "translated"
    deck_digest.digest_pptx(deck, out, max_chunk_slides=2)
    prog = json.loads((out / "progress.json").read_text(encoding="utf-8"))
    ranges = [tuple(c["slide_range"]) for c in prog["chunks"]]
    # 节1(1-1) 单独一块；节2(2-5) 超长 → 内部切成 (2-3),(4-5)
    assert ranges == [(1, 1), (2, 3), (4, 5)]


def test_section_banner_in_skeleton(deck: Path, tmp_path: Path):
    out = tmp_path / "translated"
    res = deck_digest.digest_pptx(deck, out)
    text = res.chunks[0].read_text(encoding="utf-8")
    assert "## ▸ 第一部分 理论与仿真" in text  # 节标题横幅
    assert "### Slide 003" in text  # 页降为三级标题


def test_section_starts_override(deck: Path, tmp_path: Path):
    """人工定界优先于自动检测（适用于全 deck 共用一个版式的汇报稿）。"""
    out = tmp_path / "translated"
    res = deck_digest.digest_pptx(deck, out, section_starts=[3, 5])
    assert [(s["start"], s["end"]) for s in res.sections] == [(1, 2), (3, 4), (5, 5)]


def test_title_only_slide_is_section_start(tmp_path: Path):
    """「只有标题、无正文、无图」的页应被识为分节页（即使版式普通）。"""
    prs = Presentation()
    lay = prs.slide_layouts
    s = prs.slides.add_slide(lay[0])
    s.shapes.title.text = "开场"
    s = prs.slides.add_slide(lay[1])  # 普通版式，但只填标题
    s.shapes.title.text = "第二部分 实验平台"
    s = prs.slides.add_slide(lay[1])
    s.shapes.title.text = "接线细节"
    s.placeholders[1].text_frame.text = "共 8 个通道，需要同步触发"
    deck2 = tmp_path / "deck2.pptx"
    prs.save(str(deck2))

    res = deck_digest.digest_pptx(deck2, tmp_path / "out2")
    assert [(x["start"], x["end"]) for x in res.sections] == [(1, 1), (2, 3)]
    assert res.sections[1]["title"] == "第二部分 实验平台"


def test_skeleton_keeps_image_text_binding(deck: Path, tmp_path: Path):
    """骨架必须把「图 + 位置 + 邻近文字 + 解读空位 + 合成渲染链接」写在一起。"""
    out = tmp_path / "translated"
    res = deck_digest.digest_pptx(deck, out)
    text = "\n".join(p.read_text(encoding="utf-8") for p in res.chunks)

    assert "解读：_(待填)_" in text
    assert "![](images/" in text
    assert "Fig. 1 声压幅值分布" in text  # 图注作为「邻近文字」与图绑定
    assert "重复图：同 Slide 003 图 1" in text  # 去重指向首次出现处
    assert "renders/slide-003.png" in text  # 整页合成渲染链接（待 --render 生成）
    assert "演讲者备注" in text
    assert "版面布局" in text
    assert "位置：" in text


def test_sidecar_and_progress(deck: Path, tmp_path: Path):
    out = tmp_path / "translated"
    res = deck_digest.digest_pptx(deck, out, batch_figure=5, batch_text=15)

    lines = (out / "sidecar" / "slides.jsonl").read_text(encoding="utf-8").splitlines()
    assert len(lines) == 5
    third = json.loads(lines[2])
    assert third["index"] == 3
    assert len(third["pictures"]) == 2
    assert third["pictures"][0]["img_rel"].startswith("images/")
    assert third["pictures"][0]["near_text"]  # 邻近文字已推断
    assert third["layout_map"].strip()

    prog = json.loads((out / "progress.json").read_text(encoding="utf-8"))
    assert prog["total_slides"] == 5
    assert prog["figure_slides"] == [3, 4]
    assert prog["text_slides"] == [1, 2, 5]
    assert prog["batch_policy"] == {
        "figure_slides_per_batch": 5,
        "text_slides_per_batch": 15,
    }
    assert prog["narrative_review"] == "pending"
    assert len(prog["chunks"]) == len(res.chunks)
    assert (out / "index.md").is_file()
    assert "分块导航" in (out / "index.md").read_text(encoding="utf-8")


def test_verify_links_pending_renders_not_broken(deck: Path, tmp_path: Path):
    out = tmp_path / "translated"
    deck_digest.digest_pptx(deck, out)
    rep = deck_digest.verify_links(out)
    assert rep["broken"] == []  # 图片链接全部有效
    assert len(rep["pending_renders"]) == 2  # 两个含图页的渲染待生成
    assert rep["ok"] >= 3


def test_rerun_reuses_and_preserves_interpretations(deck: Path, tmp_path: Path):
    out = tmp_path / "translated"
    deck_digest.digest_pptx(deck, out)
    # 选一个含图片（因而含「解读」空位）的分块
    part = next(
        p
        for p in sorted(out.glob("part_*.md"))
        if "解读：_(待填)_" in p.read_text(encoding="utf-8")
    )
    filled = part.read_text(encoding="utf-8").replace(
        "解读：_(待填)_", "解读：EP 附近的声压场对称性破缺", 1
    )
    part.write_text(filled, encoding="utf-8")

    res2 = deck_digest.digest_pptx(deck, out)  # 不加 force
    assert res2.reused is True
    assert "EP 附近的声压场对称性破缺" in part.read_text(encoding="utf-8")


def test_fixed_chunking(deck: Path, tmp_path: Path):
    out = tmp_path / "translated"
    res = deck_digest.digest_pptx(deck, out, chunk_by="fixed", chunk_size=2)
    assert len(res.chunks) == 3  # 5 页 → 2 + 2 + 1


def test_force_purges_stale_part_files(deck: Path, tmp_path: Path):
    """全量重建须清掉上一次残留、本次不再产出的 part_*.md（文件名随分节变化）。"""
    out = tmp_path / "translated"
    deck_digest.digest_pptx(deck, out, chunk_by="fixed", chunk_size=2)
    assert len(list(out.glob("part_*.md"))) == 3
    # 改按分节（→ 1 块）并 force：旧的 3 个应被清到只剩新的 1 个
    res = deck_digest.digest_pptx(deck, out, chunk_by="section", force=True)
    assert len(res.chunks) == 1
    assert len(list(out.glob("part_*.md"))) == 1
    # verify_links 不再因陈旧骨架重复计数（渲染待生成=含图页数）
    rep = deck_digest.verify_links(out)
    assert rep["broken"] == []
    assert len(rep["pending_renders"]) == res.n_figure_slides


def test_bad_input(tmp_path: Path):
    with pytest.raises(FileNotFoundError):
        deck_digest.digest_pptx(tmp_path / "nope.pptx", tmp_path / "out")
    legacy = tmp_path / "old.ppt"
    legacy.write_bytes(b"x")
    with pytest.raises(ValueError):
        deck_digest.digest_pptx(legacy, tmp_path / "out")


# ---------------------------------------------------------------------------
# 不可直读格式（gif 动图 / wmf 矢量）的预览
# ---------------------------------------------------------------------------
def test_even_indices():
    """抽帧下标均匀分布且必含首末帧。"""
    assert deck_digest._even_indices(1, 3) == [0]
    assert deck_digest._even_indices(5, 1) == [0]
    assert deck_digest._even_indices(3, 3) == [0, 1, 2]
    assert deck_digest._even_indices(5, 3) == [0, 2, 4]
    assert deck_digest._even_indices(5, 10) == [0, 1, 2, 3, 4]


def test_gif_preview_frames_extracted(tmp_path: Path):
    """gif 动图 LLM 不能直读 → 抽首/中/末帧为 PNG 预览，并在骨架里说明。"""
    prs = Presentation()
    lay = prs.slide_layouts
    s = prs.slides.add_slide(lay[5])
    s.shapes.title.text = "时变超表面仿真动画"
    gif5 = _gif(
        [(10, 10, 10), (60, 20, 20), (120, 40, 40), (200, 80, 80), (250, 250, 250)]
    )
    s.shapes.add_picture(io.BytesIO(gif5), Inches(1.0), Inches(1.6), Inches(4), Inches(3))
    cap = s.shapes.add_textbox(Inches(1.0), Inches(5.0), Inches(6), Inches(0.6))
    cap.text_frame.text = "动画：增益随时间演化"
    deck_gif = tmp_path / "gifdeck.pptx"
    prs.save(str(deck_gif))

    out = tmp_path / "translated"
    res = deck_digest.digest_pptx(deck_gif, out, gif_frames=3)

    # 5 帧抽 3 帧（首/中/末）落盘
    prev = list((out / "images" / "previews").glob("*.png"))
    assert len(prev) == 3
    manifest = json.loads((out / "images_manifest.json").read_text(encoding="utf-8"))
    entry = next(iter(manifest.values()))
    assert entry["ext"] == "gif"
    assert entry["frame_count"] == 5
    assert len(entry["previews"]) == 3
    # 骨架把动图说明 + 预览链接 + 邻近文字写在一起（图文不分离）
    text = res.chunks[0].read_text(encoding="utf-8")
    assert "动图" in text
    assert "images/previews/" in text
    assert "动画：增益随时间演化" in text
    # 预览帧已存在 → verify_links 无 broken（仅 renders 为 pending）
    rep = deck_digest.verify_links(out)
    assert rep["broken"] == []


def test_make_previews_vector_needs_libreoffice(tmp_path: Path):
    """矢量图（wmf）Pillow 读不了 → 给确定性预览路径，标记待 LibreOffice。"""
    img = tmp_path / "images" / "s131_p4_abcd1234.wmf"
    img.parent.mkdir(parents=True)
    img.write_bytes(b"\xd7\xcd\xc6\x9a" + b"\x00" * 16)
    info = deck_digest._make_previews(img, tmp_path, "wmf", 3)
    assert info["preview_status"] == "needs_libreoffice"
    assert info["previews"] == ["images/previews/s131_p4_abcd1234.png"]
    # 预览目录已建（等 --render 阶段填），但 PNG 尚未生成
    assert not (tmp_path / "images" / "previews" / "s131_p4_abcd1234.png").exists()


def test_make_vector_previews_deferred_without_soffice(tmp_path: Path, monkeypatch):
    """无 LibreOffice 时，矢量图预览延后（不报错、不误改 manifest 状态）。"""
    from pysci.skills.document_writing.tools.config import settings

    monkeypatch.setattr(type(settings), "find_libreoffice", lambda self: None)
    manifest = {
        "deadbeef": {
            "rel": "images/s131_p4_abcd1234.wmf",
            "ext": "wmf",
            "preview_status": "needs_libreoffice",
        }
    }
    done, note = deck_digest._make_vector_previews(tmp_path, manifest)
    assert done == 0
    assert "待 LibreOffice" in note
    assert manifest["deadbeef"]["preview_status"] == "needs_libreoffice"
