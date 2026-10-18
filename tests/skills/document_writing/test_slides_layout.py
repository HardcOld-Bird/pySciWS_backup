"""`compose slides layout`：标题对齐与分焦点块引用框规则（a2 §二.4 机械面）。"""

from __future__ import annotations

from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx", reason="需要 python-pptx")

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

from pysci.skills.document_writing.tools import compose  # noqa: E402


def _deck(tmp_path: Path, *, misaligned: bool = False) -> tuple[Path, dict[str, int]]:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])

    def rect(x, y, w, h):
        return s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
        )

    def text(x, y, w, h, body):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tb.text_frame.text = body
        return tb

    t1 = text(0.5, 1.0, 3.0, 0.44, "块一")
    t2 = text(6.5, 1.2 if misaligned else 1.0, 3.0, 0.44, "块二")
    a1, a2 = rect(0.5, 2.0, 5.0, 3.0), rect(6.5, 2.0, 5.0, 3.0)
    ref1 = text(0.5, 6.5, 5.0, 0.7, "[1] 文献甲\n[2] 文献乙")
    ref2 = text(
        6.7 if misaligned else 6.5,
        6.6 if misaligned else 6.5,
        5.0,
        0.5 if misaligned else 0.7,
        "[1] 文献丙\n[2] 文献丁\n[3] 文献戊"
        if misaligned
        else "[1] 文献丙\n[2] 文献丁",
    )
    deck = tmp_path / "layout.pptx"
    prs.save(str(deck))
    return deck, {
        "t1": t1.shape_id,
        "t2": t2.shape_id,
        "a1": a1.shape_id,
        "a2": a2.shape_id,
        "ref1": ref1.shape_id,
        "ref2": ref2.shape_id,
    }


def _run(deck: Path, ids: dict[str, int], *extra: str) -> int:
    return compose.main(
        [
            "slides",
            "layout",
            str(deck),
            "--page",
            "1",
            "--title",
            f"{ids['t1']},{ids['t2']}",
            "--refs",
            f"{ids['ref1']}:{ids['a1']}",
            "--refs",
            f"{ids['ref2']}:{ids['a2']}",
            *extra,
        ]
    )


def test_aligned_deck_passes_with_itemized_readout(tmp_path, capsys):
    deck, ids = _deck(tmp_path)
    capsys.readouterr()
    assert _run(deck, ids, "--strict") == 0
    out = capsys.readouterr().out
    assert "PASS" in out and "焦点块标题对齐" in out
    assert "Δtop=0.000in" in out and "Δleft=0.000in" in out
    assert "左缘对齐" in out and "2 条文献 / 2 行" in out
    assert "多引用框上缘一致" in out and "多引用框下缘一致" in out


def test_misaligned_deck_fails_each_rule(tmp_path, capsys):
    deck, ids = _deck(tmp_path, misaligned=True)
    capsys.readouterr()
    assert _run(deck, ids, "--strict") == 1
    out = capsys.readouterr().out
    assert "[✗] 焦点块标题对齐" in out and "Δtop=0.200in" in out
    assert "[✗] 引用框" in out and "左缘对齐" in out
    assert "3 条文献 / 3 行" in out, "≤3 条一行一文献：本例合规"
    assert "[✗] 多引用框上缘一致" in out
    assert "[✗] 多引用框下缘一致" in out


def test_ref_line_rule_flags_more_lines_than_refs(tmp_path, capsys):
    deck, ids = _deck(tmp_path)
    # 两条文献拆成三行 → 违规
    prs = Presentation(str(deck))
    for sh in prs.slides[0].shapes:
        if sh.shape_id == ids["ref1"]:
            sh.text_frame.text = "[1] 文献甲（上一行）\n（续行）\n[2] 文献乙"
    prs.save(str(deck))
    capsys.readouterr()
    assert _run(deck, ids, "--strict") == 1
    out = capsys.readouterr().out
    assert "2 条文献 / 3 行" in out and "[✗] 引用框" in out


def test_json_and_skips(tmp_path, capsys):
    import json

    deck, ids = _deck(tmp_path)
    assert compose.main(["slides", "layout", str(deck), "--page", "1", "--json"]) == 0
    data = json.loads(capsys.readouterr().out)
    assert data["violations"] == 0 and data["checks"] == []
    assert any("未给 --title" in s for s in data["skipped"])
    assert any("未给 --refs" in s for s in data["skipped"])


def test_unknown_shape_and_page_are_rejected(tmp_path, capsys):
    deck, ids = _deck(tmp_path)
    assert (
        compose.main(
            ["slides", "layout", str(deck), "--page", "1", "--title", "999999"]
        )
        == 2
    )
    assert "形状 ID 不存在" in capsys.readouterr().err
    assert compose.main(["slides", "layout", str(deck), "--page", "9"]) == 2
