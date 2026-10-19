"""`compose slides audit` 与字体全槽 API：正反例（装饰/ea/图槽/公式/引用框/动画）。"""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx", reason="需要 python-pptx")
Image = pytest.importorskip("PIL.Image", reason="需要 Pillow")
yaml = pytest.importorskip("yaml", reason="需要 PyYAML")

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches  # noqa: E402

from pysci.skills.document_writing.tools import (  # noqa: E402
    compose,
    slides_audit,
    slides_fonts,
)

_P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_TIMING = (
    f'<p:timing xmlns:p="{_P_NS}"><p:tnLst><p:par><p:cTn id="1"><p:childTnLst>'
    '<p:seq nextAc="seek"><p:cTn id="2"><p:childTnLst><p:par><p:cTn id="3">'
    '<p:stCondLst><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>'
    "</p:stCondLst></p:cTn></p:par></p:childTnLst></p:cTn></p:seq>"
    "</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"
)


def _png(path: Path, color=(10, 20, 30)) -> Path:
    Image.new("RGB", (64, 48), color).save(path)
    return path


def _set_dist(shape) -> None:
    for p in shape.text_frame.paragraphs:
        p.alignment = PP_ALIGN.DISTRIBUTE


def _deck(tmp_path: Path, *, break_decoration=False, break_ea=False, wrong_pic=False):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    card = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.1), Inches(0.1), Inches(13.0), Inches(7.3)
    )
    card.name = "白卡"
    title = s.shapes.add_textbox(Inches(0.2), Inches(0.2), Inches(1.91), Inches(0.572))
    title.text_frame.text = "研究背景"
    _set_dist(title)
    cap = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(1.2), Inches(2.0175), Inches(0.438)
    )
    cap.name = "胶囊"
    cap.text_frame.text = "PT对称"
    _set_dist(cap)
    body = s.shapes.add_textbox(Inches(0.2), Inches(2.0), Inches(5.0), Inches(1.0))
    body.text_frame.text = "正文一行"
    ref = s.shapes.add_textbox(Inches(0.1), Inches(6.5), Inches(5.0), Inches(0.6))
    ref.text_frame.text = "[1] 文献甲\n[2] 文献乙"
    _png(tmp_path / "slot.png")
    if wrong_pic:
        _png(tmp_path / "other.png", (200, 0, 0))
    pic = s.shapes.add_picture(
        str(tmp_path / ("other.png" if wrong_pic else "slot.png")),
        Inches(6.0),
        Inches(2.0),
        Inches(3.0),
        Inches(2.0),
    )
    for sh in (title, cap, body, ref):
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                slides_fonts.set_font_slots(
                    r,
                    latin="汉仪大宋简",
                    ea="汉仪大宋简",
                    cs="汉仪大宋简",
                    panose="02010600030101010101",
                )
    if break_ea:
        r = title.text_frame.paragraphs[0].runs[0]
        ea = r._r.get_or_add_rPr().find(qn("a:ea"))
        if ea is not None:
            r._r.get_or_add_rPr().remove(ea)
    if break_decoration:
        card._element.getparent().remove(card._element)
    s.element.append(slides_audit_yaml_element())
    deck = tmp_path / f"deck-{break_decoration}-{break_ea}-{wrong_pic}.pptx"
    prs.save(str(deck))
    return deck, {
        "card": card.shape_id if not break_decoration else None,
        "title": title.shape_id,
        "cap": cap.shape_id,
        "body": body.shape_id,
        "ref": ref.shape_id,
        "pic": pic.shape_id,
    }


def slides_audit_yaml_element():
    from lxml import etree

    return etree.fromstring(_TIMING)


def _spec(tmp_path: Path, ids: dict) -> Path:
    md5 = hashlib.md5((tmp_path / "slot.png").read_bytes()).hexdigest()
    doc = {
        "deck": "test",
        "pages": [
            {
                "page": 1,
                "decoration": [
                    {"name": "白卡", "box_in": [0.1, 0.1, 13.1, 7.4]},
                ],
                "fonts": [
                    {"role": "标题", "shapes": [ids["title"]]},
                    {"role": "胶囊", "shapes": [ids["cap"]]},
                    {"role": "正文", "shapes": [ids["body"], ids["ref"]]},
                ],
                "slots": [
                    {
                        "shape": ids["pic"],
                        "figure_no": "Fig. 7",
                        "anchor": "RPP 86 Fig.7 场分布",
                        "md5": md5,
                        "px": [64, 48],
                        "source": "slot.png",
                    }
                ],
                "titles": [{"shape": ids["title"]}],
                "capsules": [{"shape": ids["cap"], "cite_chars": 3}],
                "refs": [{"ref": ids["ref"], "anchor": ids["card"] or 2}],
                "animation": True,
                "manual": ["装饰质感（目检）"],
            }
        ],
    }
    p = tmp_path / "spec.yaml"
    p.write_text(yaml.safe_dump(doc, allow_unicode=True), encoding="utf-8")
    return p


def test_font_api_writes_all_slots_and_panose():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    tb.text_frame.text = "模板"
    src = tb.text_frame.paragraphs[0].runs[0]
    slides_fonts.set_font_slots(src, latin="L1", ea="E1", cs="C1", panose="P1")
    tb2 = s.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1))
    tb2.text_frame.text = "目标"
    dst = tb2.text_frame.paragraphs[0].runs[0]
    dst.font.name = "only-latin"
    slides_fonts.copy_font_slots(src, dst)
    slots = slides_fonts.font_slots(dst)
    assert slots["latin"]["typeface"] == "L1"
    assert slots["ea"]["typeface"] == "E1" and slots["cs"]["typeface"] == "C1"
    assert slots["ea"]["panose"] == "P1"
    assert slides_fonts.slots_complete(slots)


def test_audit_positive_deck_passes(tmp_path):
    deck, ids = _deck(tmp_path)
    rep = slides_audit.audit_deck(deck, _spec(tmp_path, ids))
    assert rep.ok, rep.summary()
    assert rep.manual and "目检" in rep.manual[0].detail
    assert any(i.name.startswith("图槽") and i.status == "pass" for i in rep.items)


def test_audit_negative_decoration_missing(tmp_path):
    deck, ids = _deck(tmp_path, break_decoration=True)
    rep = slides_audit.audit_deck(deck, _spec(tmp_path, ids))
    assert not rep.ok and any("未找到匹配形状" in i.detail for i in rep.failures)


def test_audit_negative_missing_ea(tmp_path):
    deck, ids = _deck(tmp_path, break_ea=True)
    rep = slides_audit.audit_deck(deck, _spec(tmp_path, ids))
    assert any("缺槽" in i.detail and "ea" in i.detail for i in rep.failures), (
        rep.summary()
    )


def test_audit_negative_slot_md5_mismatch(tmp_path):
    deck, ids = _deck(tmp_path, wrong_pic=True)
    rep = slides_audit.audit_deck(deck, _spec(tmp_path, ids))
    assert any("md5" in i.detail and "≠" in i.detail for i in rep.failures), (
        rep.summary()
    )


def test_audit_cli_strict_and_json(tmp_path, capsys):
    deck, ids = _deck(tmp_path)
    spec = _spec(tmp_path, ids)
    assert (
        compose.main(["slides", "audit", str(deck), "--spec", str(spec), "--strict"])
        == 0
    )
    assert "PASS" in capsys.readouterr().out
    bad, bad_ids = _deck(tmp_path, break_ea=True)
    assert (
        compose.main(["slides", "audit", str(bad), "--spec", str(spec), "--json"]) == 0
    )
    import json

    data = json.loads(capsys.readouterr().out)
    assert data["failures"] >= 1 and data["ok"] is False


# ---------------------------------------------------------------------------
# A/B/C：绑定一致性 / 位置与文本判据 / 注记新鲜度（L.l3 三工单）
# ---------------------------------------------------------------------------
def _box_in(shape) -> list[float]:
    return [
        round(shape.left / 914400, 3),
        round(shape.top / 914400, 3),
        round((shape.left + shape.width) / 914400, 3),
        round((shape.top + shape.height) / 914400, 3),
    ]


def _mutated_spec(tmp_path, ids, mutate) -> Path:
    doc = yaml.safe_load(_spec(tmp_path, ids).read_text(encoding="utf-8"))
    mutate(doc["pages"][0], ids)
    p = tmp_path / "spec-mut.yaml"
    p.write_text(yaml.safe_dump(doc, allow_unicode=True), encoding="utf-8")
    return p


def test_duplicate_placeholder_binding_fails(tmp_path):
    deck, ids = _deck(tmp_path)
    page = {
        "page": 1,
        "placeholders": [{"md5": "x"}],
        "slots": [
            {"shape": ids["title"], "placeholder": {"md5": "x"}},
            {"shape": ids["cap"], "placeholder": {"md5": "x"}},
        ],
    }
    p = tmp_path / "spec-dup.yaml"
    p.write_text(yaml.safe_dump({"pages": [page]}), encoding="utf-8")
    rep = slides_audit.audit_deck(deck, p)
    assert any(
        i.name.startswith("绑定 p1") and i.status == "fail" for i in rep.items
    ), rep.summary()


def test_semantic_source_anchor_mismatch_is_manual(tmp_path):
    deck, ids = _deck(tmp_path)

    def mutate(page, ids):
        page["slots"][0]["anchor"] = "PRL 123 Fig.2 三面板"
        page["slots"][0]["source"] = (
            "文献图片/第5页-异常点-NatMater18-783-2019/fig7.png"
        )

    # 该来源不存在于 tmp_path 相对面 → 会同时报来源缺失；语义提示先于内容检查
    rep = slides_audit.audit_deck(deck, _mutated_spec(tmp_path, ids, mutate))
    assert any(
        "绑定存疑" in i.name and "NatMater" in i.detail and "PRL" in i.detail
        for i in rep.items
    )


def test_cyclic_rotation_is_suspected_and_ack_silences(tmp_path):
    deck, ids = _deck(tmp_path)

    def page_doc(ack: bool) -> dict:
        slots = [
            {"shape": ids["title"], "placeholder": {"md5": "b"}},
            {"shape": ids["cap"], "placeholder": {"md5": "c"}},
            {"shape": ids["body"], "placeholder": {"md5": "a"}},
        ]
        p = {
            "page": 1,
            "placeholders": [{"md5": "a"}, {"md5": "b"}, {"md5": "c"}],
            "slots": slots,
        }
        if ack:
            p["binding_ack"] = True
        return p

    for ack, expect in ((False, True), (True, False)):
        p = tmp_path / f"spec-rot-{ack}.yaml"
        p.write_text(yaml.safe_dump({"pages": [page_doc(ack)]}), encoding="utf-8")
        rep = slides_audit.audit_deck(deck, p)
        got = any("绑定错位疑似" in i.name for i in rep.items)
        assert got is expect, rep.summary()


def test_expect_box_pass_and_fail(tmp_path):
    deck, ids = _deck(tmp_path)
    prs = Presentation(str(deck))
    pic = next(sh for sh in prs.slides[0].shapes if sh.shape_id == ids["pic"])
    good = _box_in(pic)

    def mutate(page, ids):
        page["slots"][0]["expect_box"] = good
        page["slots"][0]["box_tol_in"] = 0.05

    rep = slides_audit.audit_deck(deck, _mutated_spec(tmp_path, ids, mutate))
    assert any(i.name.startswith("位置") and i.status == "pass" for i in rep.items)

    def mutate_bad(page, ids):
        page["slots"][0]["expect_box"] = [0.0, 0.0, 1.0, 1.0]

    rep_bad = slides_audit.audit_deck(deck, _mutated_spec(tmp_path, ids, mutate_bad))
    assert any(i.name.startswith("位置") and i.status == "fail" for i in rep_bad.items)


def test_expect_text_pass_and_fail(tmp_path):
    deck, ids = _deck(tmp_path)

    def mutate_ok(page, ids):
        page["slots"].append({"shape": ids["title"], "expect_text": "研究"})

    rep = slides_audit.audit_deck(deck, _mutated_spec(tmp_path, ids, mutate_ok))
    assert any(i.name.startswith("文本") and i.status == "pass" for i in rep.items)

    def mutate_bad(page, ids):
        page["slots"].append({"shape": ids["title"], "expect_text": "不存在的文字"})

    rep_bad = slides_audit.audit_deck(deck, _mutated_spec(tmp_path, ids, mutate_bad))
    assert any(i.name.startswith("文本") and i.status == "fail" for i in rep_bad.items)


def test_stale_note_hint_when_md5_matches(tmp_path):
    deck, ids = _deck(tmp_path)

    def mutate(page, ids):
        page["slots"][0]["note"] = "批语：待换"

    rep = slides_audit.audit_deck(deck, _mutated_spec(tmp_path, ids, mutate))
    assert any("注记过期" in i.name and i.status == "manual" for i in rep.items), (
        rep.summary()
    )
    assert rep.ok, "注记过期只是提示，不计 FAIL"
