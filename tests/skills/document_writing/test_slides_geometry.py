"""`compose slides geometry` 的几何体检回归（缺陷 1：遮挡/越界/图片互叠的机械原料）。

渲染闸只防「没渲染 / 渲染成空白」，命中不了「渲染出来了但版面是坏的」——本文件钉住几何端
的三类硬判据：外框边被更高 z 序图片**覆盖**、图元**越界**（画布/外框）、**图片互叠**面积。
用临时构造的迷你 pptx 制造可控缺陷，断言三条各有命中，且该命令**只读不写盘**。
"""

from __future__ import annotations

from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx", reason="需要 python-pptx")
Image = pytest.importorskip("PIL.Image", reason="需要 Pillow")

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

from pysci.skills.document_writing.tools import compose, slides_visual  # noqa: E402


def _png(path: Path, color: tuple[int, int, int] = (30, 90, 200)) -> Path:
    Image.new("RGB", (60, 40), color).save(path)
    return path


def _defective_deck(tmp_path: Path) -> Path:
    """构造三类缺陷各一的迷你 deck：外框、压外框上边的图、互叠的两图、越出画布的矩形。"""
    prs = Presentation()  # 默认 10x7.5 in 画布
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(8), Inches(5)
    )
    frame.name = "外框测试"
    img = _png(tmp_path / "shot.png")
    # 压住外框上边（x 6–8 段）
    slide.shapes.add_picture(str(img), Inches(6), Inches(0.5), Inches(2), Inches(2))
    # 两张互叠（交叠 1.5x1.5 = 2.25 in²）
    slide.shapes.add_picture(str(img), Inches(2), Inches(2), Inches(3), Inches(2))
    slide.shapes.add_picture(str(img), Inches(3.5), Inches(2.5), Inches(3), Inches(2))
    # 越出画布右侧
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(9), Inches(6), Inches(3), Inches(3)
    )
    deck = tmp_path / "defective.pptx"
    prs.save(str(deck))
    return deck


def _clean_deck(tmp_path: Path) -> Path:
    """无缺陷对照：只有一个居中矩形 + 一张位于其内、带黑描边的图。"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(8), Inches(5)
    )
    pic = slide.shapes.add_picture(
        str(_png(tmp_path / "ok.png")), Inches(2), Inches(2), Inches(2), Inches(2)
    )
    pic.line.color.rgb = RGBColor(0, 0, 0)  # 预设第 4 条：非整页图加黑描边
    deck = tmp_path / "clean.pptx"
    prs.save(str(deck))
    return deck


def test_geometry_reports_edge_coverage_overlap_and_overflow(tmp_path):
    rep = slides_visual.geometry_report(_defective_deck(tmp_path))
    sg = rep.slides[0]
    assert sg.frame is not None and sg.frame_source == "auto"
    assert sg.frame.name == "外框测试"
    covered = [e for e in sg.edges if e.frac > 0]
    assert any(e.edge == "上边" and e.frac > 0.2 for e in covered), sg.edges
    assert sg.overlaps, "互叠的两张图未被报出"
    assert sg.off_canvas, "越出画布的矩形未被报出"
    assert rep.violations > 0 and not rep.ok


def test_clean_deck_has_no_violations(tmp_path):
    rep = slides_visual.geometry_report(_clean_deck(tmp_path))
    assert rep.ok and rep.violations == 0, rep.summary()


def test_explicit_frame_name_selects_the_frame(tmp_path):
    deck = _defective_deck(tmp_path)
    rep = slides_visual.geometry_report(deck, frame="外框测试")
    assert rep.slides[0].frame_source == "explicit"
    # 名字不存在时不臆造外框（只报越界/互叠）
    rep2 = slides_visual.geometry_report(deck, frame="不存在的外框")
    assert rep2.slides[0].frame is None and rep2.slides[0].edges == []


def test_off_frame_only_judges_pictures(tmp_path):
    """外框之外的标题文本框不算「越出外框」（原型期实测的假阳性面）。"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(8), Inches(5)
    )
    tb = slide.shapes.add_textbox(Inches(1), Inches(0.2), Inches(4), Inches(0.5))
    tb.text_frame.text = "标题（在外框上方，属设计）"
    deck = tmp_path / "titled.pptx"
    prs.save(str(deck))
    rep = slides_visual.geometry_report(deck)
    assert rep.slides[0].off_frame == [], rep.slides[0].off_frame


def test_cli_geometry_reports_only_by_default_and_strict_fails(tmp_path, capsys):
    deck = _defective_deck(tmp_path)
    assert compose.main(["slides", "geometry", str(deck)]) == 0  # 默认只报告
    out = capsys.readouterr().out
    assert "外框" in out and "违规" in out
    assert compose.main(["slides", "geometry", str(deck), "--strict"]) == 1
    assert (
        compose.main(["slides", "geometry", str(_clean_deck(tmp_path)), "--strict"])
        == 0
    )


# ---------------------------------------------------------------------------
# backlog 20261019-050008：近满幅白卡（>95%）不得回退到焦点块面板
# ---------------------------------------------------------------------------
def _near_full_card_deck(tmp_path: Path) -> Path:
    """10×7.5 画布：100% 背景 + 97.7% 白卡 + 两个面板 + 右侧图（在白卡内、面板外）。"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(7.5)
    )
    bg.name = "背景"
    card = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.05), Inches(0.05), Inches(9.9), Inches(7.4)
    )
    card.name = "白卡"
    left = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(0.2), Inches(4.0), Inches(6.5)
    )
    left.name = "面板左"
    right = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(0.2), Inches(4.0), Inches(6.5)
    )
    right.name = "面板右"
    # 右缘图落在白卡内（x 8.5–9.5 < 白卡 9.95），但在两个面板之外
    slide.shapes.add_picture(
        str(_png(tmp_path / "right.png")),
        Inches(8.5),
        Inches(2),
        Inches(1),
        Inches(1),
    )
    deck = tmp_path / "near-full-card.pptx"
    prs.save(str(deck))
    return deck


def test_near_full_page_card_is_detected_as_frame(tmp_path):
    deck = _near_full_card_deck(tmp_path)
    rep = slides_visual.geometry_report(deck)
    sg = rep.slides[0]
    assert sg.frame is not None and sg.frame.name == "白卡", sg.frame
    assert sg.frame_source == "auto"
    assert not sg.frame_note
    assert sg.off_frame == [], sg.off_frame


def test_auto_and_explicit_white_card_agree_on_off_frame(tmp_path):
    deck = _near_full_card_deck(tmp_path)
    auto = slides_visual.geometry_report(deck).slides[0]
    explicit = slides_visual.geometry_report(deck, frame="白卡").slides[0]
    assert explicit.frame_source == "explicit"
    assert len(auto.off_frame) == len(explicit.off_frame) == 0


def test_uncertain_auto_frame_warns_instead_of_off_frame(tmp_path):
    """小矩形占多数时 auto 仍选中 32% 卡，但给存疑提示并抑制越界行。"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    card = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(0.2), Inches(4.0), Inches(6.0)
    )
    card.name = "小卡"
    for i, x in enumerate((4.8, 7.0)):
        sh = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(1), Inches(3), Inches(3)
        )
        sh.name = f"面板{i + 1}"
    slide.shapes.add_picture(
        str(_png(tmp_path / "outside.png")),
        Inches(9.0),
        Inches(6.0),
        Inches(0.8),
        Inches(0.8),
    )
    deck = tmp_path / "uncertain.pptx"
    prs.save(str(deck))
    rep = slides_visual.geometry_report(deck)
    sg = rep.slides[0]
    assert sg.frame is not None and sg.frame.name == "小卡"
    assert "存疑" in sg.frame_note, sg.frame_note
    assert sg.off_frame == [], "存疑时不得直接输出越界违规"
    assert sg.edges == []
    assert "外框识别存疑" in rep.summary()
    assert rep.as_dict()["slides"][0]["frame_note"]


def test_cli_geometry_json_is_machine_readable(tmp_path, capsys):
    import json

    deck = _defective_deck(tmp_path)
    assert compose.main(["slides", "geometry", str(deck), "--json"]) == 0
    data = json.loads(capsys.readouterr().out)
    assert data["pptx"].endswith("defective.pptx")
    assert data["violations"] > 0 and data["slides"][0]["frame"]["name"] == "外框测试"


def test_geometry_is_read_only(tmp_path):
    """只读纪律：几何体检**不写盘**（受审目录只读时也照跑）。"""
    deck = _defective_deck(tmp_path)
    before = {p.name for p in tmp_path.iterdir()}
    slides_visual.geometry_report(deck)
    compose.main(["slides", "geometry", str(deck), "--json"])
    assert {p.name for p in tmp_path.iterdir()} == before


def test_render_prints_ready_to_paste_verify_command(tmp_path, capsys):
    """缺陷 2 的用法面：渲染后打印可直接粘贴的 verify 命令（不必记 --render-dir 旗标名）。"""
    deck = _clean_deck(tmp_path)
    rc = compose.main(["slides", "render", str(deck), "--out", str(tmp_path / "r")])
    if rc == 3:  # 无 LibreOffice：渲染链不可用，跳过该断言
        pytest.skip("未检测到 LibreOffice")
    out = capsys.readouterr().out
    assert "--render-dir" in out and "slides verify" in out


# ---------------------------------------------------------------------------
# 连接符越界：bbox 不可信 → 单列提示、不计违规（backlog 20261016-033419-writing ②）
# ---------------------------------------------------------------------------
def _connector_deck(tmp_path: Path, *, overflow: bool) -> Path:
    prs = Presentation()  # 默认 10x7.5 in 画布
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if overflow:
        # 肘形连接符越出画布右下（PowerPoint 给连接符存的 bbox 常远大于实际路径）
        slide.shapes.add_connector(
            MSO_CONNECTOR.ELBOW, Inches(9), Inches(6), Inches(11), Inches(8)
        )
    else:
        slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(1), Inches(3), Inches(3)
        )
    deck = tmp_path / "connector.pptx"
    prs.save(str(deck))
    return deck


def test_connector_kind_is_detected(tmp_path):
    rep = slides_visual.geometry_report(_connector_deck(tmp_path, overflow=False))
    kinds = [b.kind for b in rep.slides[0].boxes]
    assert "connector" in kinds


def test_connector_overflow_is_hint_not_violation(tmp_path):
    """连接符越界只作提示：不进 off_canvas、不计入 violations → 页判 PASS。"""
    rep = slides_visual.geometry_report(_connector_deck(tmp_path, overflow=True))
    sg = rep.slides[0]
    assert sg.connector_hints, "连接符越界未报为提示"
    assert sg.off_canvas == [], "连接符被误计入越界（bbox 假阳性）"
    assert sg.violations == 0 and rep.ok, rep.summary()
    assert "非违规" in sg.connector_hints[0]


def test_connector_hint_does_not_mask_a_real_overflow(tmp_path):
    """同一页里，连接符提示与真正越界并存——真越界照报、照计违规。"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_connector(
        MSO_CONNECTOR.ELBOW, Inches(9), Inches(6), Inches(11), Inches(8)
    )
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(9), Inches(6), Inches(3), Inches(3)
    )  # 真越界矩形
    deck = tmp_path / "mixed.pptx"
    prs.save(str(deck))
    sg = slides_visual.geometry_report(deck).slides[0]
    assert sg.connector_hints and sg.off_canvas, (sg.connector_hints, sg.off_canvas)
    assert sg.violations > 0


def test_strict_passes_on_connector_hint_only(tmp_path):
    deck = _connector_deck(tmp_path, overflow=True)
    assert compose.main(["slides", "geometry", str(deck), "--strict"]) == 0


def test_geometry_json_exposes_connector_hints(tmp_path, capsys):
    import json

    deck = _connector_deck(tmp_path, overflow=True)
    assert compose.main(["slides", "geometry", str(deck), "--json"]) == 0
    data = json.loads(capsys.readouterr().out)
    assert data["slides"][0]["connector_hints"]


def test_visual_qa_documents_connector_hint():
    """用法文档面写明连接符提示口径（只打真本；部署副本在 worktree 里不存在）。"""
    from pysci.paths import ORCHESTRATION_ROOT

    md = (
        ORCHESTRATION_ROOT
        / "skills"
        / "document-writing"
        / "references"
        / "visual-qa.md"
    ).read_text(encoding="utf-8")
    assert "连接符" in md and "connector_hints" in md


# ---------------------------------------------------------------------------
# 透明装饰水印（学院正式预设）：不计内容图互叠/描边（backlog 20261018-223051-writing）
# ---------------------------------------------------------------------------
def _set_blip_alpha(shape, amt: int) -> None:
    """在图片 blip 上写 ``<a:alphaModFix amt=…>``（100000 = 100%）。"""
    from lxml import etree
    from pptx.oxml.ns import qn

    blip = shape._element.find(f".//{qn('a:blip')}")
    etree.SubElement(blip, qn("a:alphaModFix"), amt=str(amt))


def _watermark_deck(tmp_path: Path, *, alpha: int | None = 8000) -> Path:
    """白卡外框 + 两张带描边内容图 + 一张压住内容图、无描边的透明水印。"""
    prs = Presentation()  # 默认 10x7.5 in 画布
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(8), Inches(5)
    )
    img = _png(tmp_path / "content.png")
    for x in (1.5, 5.0):  # 两张内容图彼此不叠
        pic = slide.shapes.add_picture(
            str(img), Inches(x), Inches(2), Inches(3), Inches(2)
        )
        pic.line.color.rgb = RGBColor(0, 0, 0)
    wm = slide.shapes.add_picture(  # Picture 3：装饰水印，无 <a:ln>
        str(img), Inches(1.2), Inches(1.5), Inches(3), Inches(4)
    )
    if alpha is not None:
        _set_blip_alpha(wm, alpha)
    deck = tmp_path / f"watermark-{alpha}.pptx"
    prs.save(str(deck))
    return deck


def test_decorative_watermark_exempt_from_overlap_and_border(tmp_path):
    """8% 透明水印压住内容图且无描边，不得计入图片互叠/描边——只作读数列出。"""
    rep = slides_visual.geometry_report(_watermark_deck(tmp_path, alpha=8000))
    sg = rep.slides[0]
    assert len(sg.decorative_pics) == 1 and sg.decorative_pics[0].startswith(
        "Picture"
    ), sg.decorative_pics
    assert sg.overlaps == [], f"水印被计入图片互叠：{sg.overlaps}"
    assert sg.picture_border_issues == [], sg.picture_border_issues
    assert sg.violations == 0 and rep.ok, rep.summary()
    assert "装饰水印" in rep.summary()


@pytest.mark.parametrize("alpha", [20000, None])
def test_non_decorative_picture_still_flagged(tmp_path, alpha):
    """对照：20% 不透明或未设透明度的图片仍是内容图——互叠与缺描边照报。"""
    rep = slides_visual.geometry_report(_watermark_deck(tmp_path, alpha=alpha))
    sg = rep.slides[0]
    assert sg.decorative_pics == []
    assert sg.overlaps, "内容图互叠未报出"
    assert sg.picture_border_issues, "内容图缺描边未报出"
    assert sg.violations > 0


def test_geometry_json_exposes_decorative_pics(tmp_path, capsys):
    import json

    deck = _watermark_deck(tmp_path, alpha=8000)
    assert compose.main(["slides", "geometry", str(deck), "--json"]) == 0
    data = json.loads(capsys.readouterr().out)
    assert len(data["slides"][0]["decorative_pics"]) == 1
    assert data["violations"] == 0


def test_strict_passes_with_decorative_watermark(tmp_path):
    deck = _watermark_deck(tmp_path, alpha=8000)
    assert compose.main(["slides", "geometry", str(deck), "--strict"]) == 0


# ---------------------------------------------------------------------------
# W1 形状缺 <a:ext> 机械闸（tb-20261017-044633-4-1016ppt-1）
# ---------------------------------------------------------------------------
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _titled_deck(tmp_path: Path, *, slide_titles: list[str], mutate=None) -> Path:
    """建「仅标题」版式 deck：先装配（含可选 ``mutate(prs, slides)`` 改动）**再存盘**。

    先给标题显式写 w/h——python-pptx 的占位符默认**不写** ``<a:xfrm><a:ext>``，不逼出它
    就没有「xfrm 在、ext 缺」这一真实缺陷签名可删。标题按预设方案写成「·标题·」且框宽
    随文本估算自适应（否则会被新增的预设三项检查判违规，干扰 W1/W2 的判据）。
    """
    from pptx.util import Pt

    prs = Presentation()
    for t in slide_titles:
        s = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
        title = s.shapes.title
        title.text = f"·{t}·"
        for r in title.text_frame.paragraphs[0].runs:
            r.font.size = Pt(20)
        title.width = Inches(slides_visual.estimate_title_width_in(title))
        title.height = Inches(0.44)
    slides = list(prs.slides)
    if mutate is not None:
        mutate(prs, slides)
    deck = tmp_path / "titled.pptx"
    prs.save(str(deck))
    return deck


def _strip_ext(shape) -> None:
    """删掉形状 ``<a:xfrm>`` 里的 ``<a:ext>``（显式尺寸缺失的形态）。"""
    xfrm = shape._element.find(f".//{{{_A_NS}}}xfrm")
    xfrm.remove(xfrm.find(f"{{{_A_NS}}}ext"))


def test_w1_missing_ext_is_flagged_per_page(tmp_path):
    """缺 <a:ext> 的形状必须被逐页列出，且计入 violations（一票否决）。"""
    deck = _titled_deck(
        tmp_path,
        slide_titles=["一", "二", "三"],
        mutate=lambda prs, slides: _strip_ext(slides[1].shapes.title),
    )
    rep = slides_visual.geometry_report(deck)
    assert rep.ext_missing_total == 1
    assert rep.slides[1].missing_ext and "二" in rep.slides[1].missing_ext[0]
    assert rep.slides[0].missing_ext == []
    assert rep.violations >= 1 and not rep.ok


def test_w1_ext_present_is_clean(tmp_path):
    """python-pptx 写入的标题带显式 ext → 不误报；clean deck 仍 PASS。"""
    deck = _titled_deck(tmp_path, slide_titles=["一", "二"])
    rep = slides_visual.geometry_report(deck)
    assert rep.ext_missing_total == 0
    assert rep.ok and rep.violations == 0, rep.summary()


def test_w4_risk_line_names_render_vs_wps(tmp_path):
    """W4：缺 ext 时输出风险行「渲染图正常 ≠ WPS 正常」。"""
    deck = _titled_deck(
        tmp_path,
        slide_titles=["一", "二"],
        mutate=lambda prs, slides: _strip_ext(slides[0].shapes.title),
    )
    out = slides_visual.geometry_report(deck).summary()
    assert "渲染图正常" in out and "WPS 正常" in out


def test_w1_cli_strict_fails_on_missing_ext(tmp_path):
    deck = _titled_deck(
        tmp_path,
        slide_titles=["一", "二"],
        mutate=lambda prs, slides: _strip_ext(slides[0].shapes.title),
    )
    assert compose.main(["slides", "geometry", str(deck), "--strict"]) == 1


# ---------------------------------------------------------------------------
# W2 跨页标题一致性
# ---------------------------------------------------------------------------
def test_w2_consistent_titles_have_no_diff(tmp_path):
    deck = _titled_deck(tmp_path, slide_titles=["一", "二", "三"])
    assert slides_visual.geometry_report(deck).title_diffs == []


def test_w2_one_changed_title_is_reported(tmp_path):
    """人为改一页标题（字号）即报差异，并点名页码与差异项。"""
    from pptx.util import Pt

    def _change(prs, slides):
        run = slides[2].shapes.title.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(40)

    deck = _titled_deck(tmp_path, slide_titles=["一", "二", "三"], mutate=_change)
    diffs = slides_visual.geometry_report(deck).title_diffs
    assert len(diffs) == 1 and "第 3 页" in diffs[0] and "字号" in diffs[0], diffs


def test_w2_multi_run_same_format_is_consistent(tmp_path):
    """回归（backlog 20261017-101130）：同格式被拆成**多 run** 不得误报。

    字号/字体按 **set 去重**比对——「·」+「pySci·」两 run（rPr 逐字相同）视觉等价于单 run。
    """
    from pptx.util import Pt

    def _mutate(prs, slides):
        for s in slides:  # 全部标题显式同格式
            for r in s.shapes.title.text_frame.paragraphs[0].runs:
                r.font.size = Pt(20)
                r.font.name = "Heiti SC Light"
        para = slides[2].shapes.title.text_frame.paragraphs[0]
        text = para.text
        para.runs[0].text = text[:1]
        extra = para.add_run()
        extra.text = text[1:]
        for r in para.runs:  # 两 run 的 rPr 逐字相同
            r.font.size = Pt(20)
            r.font.name = "Heiti SC Light"

    deck = _titled_deck(tmp_path, slide_titles=["一", "二", "三"], mutate=_mutate)
    assert slides_visual.geometry_report(deck).title_diffs == []


# ---------------------------------------------------------------------------
# W3/W4 文档面
# ---------------------------------------------------------------------------
def test_w3_endpage_convention_documented_both_sides():
    from pysci.paths import ORCHESTRATION_ROOT

    wch = (
        ORCHESTRATION_ROOT / "pods" / "writing" / ".qoder" / "rules" / "charter.md"
    ).read_text(encoding="utf-8")
    rubric = (
        ORCHESTRATION_ROOT / "pods" / "reviewer" / "rubrics" / "presentation.md"
    ).read_text(encoding="utf-8")
    assert "结尾页" in wch and "感谢" in wch, "writing charter 缺结尾页惯例"
    assert "结尾页" in rubric and "扣分" in rubric, "reviewer 量表缺无结尾页扣分项"


def test_w4_wps_feasibility_recorded():
    from pysci.paths import ORCHESTRATION_ROOT

    md = (
        ORCHESTRATION_ROOT
        / "skills"
        / "document-writing"
        / "references"
        / "visual-qa.md"
    ).read_text(encoding="utf-8")
    assert "WPS" in md and "不可行" in md and "wpscli" in md, "缺 WPS 渲染可行性结论"
