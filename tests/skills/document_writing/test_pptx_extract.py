"""document_writing PPTX 提取测试。

需要一个样例 .pptx；用 python-pptx 现场构建（含标题/多级项目符号/表格/图片/演讲者备注），
再分别用 pptx_io（结构化）与 extract（markitdown）验证提取。

未安装 [writing] extra 时整体跳过（pytest.importorskip）。
"""

from __future__ import annotations

from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx", reason="需要 python-pptx：uv sync --extra writing")

from pptx.util import Inches  # noqa: E402

from pysci.skills.document_writing.tools import extract, pptx_io  # noqa: E402


def _make_png(path: Path) -> None:
    """用 Pillow（python-pptx 的依赖）生成一张 16x16 纯色 PNG。"""
    from PIL import Image

    Image.new("RGB", (16, 16), (200, 30, 30)).save(path)


def _build_sample_deck(path: Path, img: Path) -> None:
    from pptx import Presentation

    prs = Presentation()

    # 第 1 页：标题页 + 备注
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "增益超表面中的例外点"
    s1.placeholders[1].text = "2026 年度汇报"
    s1.notes_slide.notes_text_frame.text = "开场：介绍非厄米物理背景"

    # 第 2 页：多级项目符号 + 表格 + 图片 + 备注
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    tb = s2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    tf = tb.text_frame
    tf.text = "关键结果"
    p1 = tf.add_paragraph()
    p1.text = "观测到 EP 点"
    p1.level = 1
    p2 = tf.add_paragraph()
    p2.text = "拓扑保护边界态"
    p2.level = 1

    tbl_shape = s2.shapes.add_table(
        3, 2, Inches(0.5), Inches(2), Inches(4), Inches(1)
    )
    tbl = tbl_shape.table
    tbl.cell(0, 0).text = "参数"
    tbl.cell(0, 1).text = "值"
    tbl.cell(1, 0).text = "频率"
    tbl.cell(1, 1).text = "40 kHz"
    tbl.cell(2, 0).text = "增益"
    tbl.cell(2, 1).text = "0.3"

    s2.shapes.add_picture(str(img), Inches(5), Inches(2), Inches(1.5), Inches(1.5))
    s2.notes_slide.notes_text_frame.text = "强调 EP 的实验证据与理论吻合"

    prs.save(str(path))


@pytest.fixture()
def sample_pptx(tmp_path: Path) -> Path:
    img = tmp_path / "fig.png"
    _make_png(img)
    deck = tmp_path / "sample_report.pptx"
    _build_sample_deck(deck, img)
    return deck


def test_read_pptx_structure(sample_pptx: Path) -> None:
    slides = pptx_io.read_pptx(sample_pptx)
    assert len(slides) == 2

    s1 = slides[0]
    assert "例外点" in s1.title
    assert "非厄米" in s1.notes

    s2 = slides[1]
    # 多级项目符号：至少含两条正文
    joined = "\n".join(s2.paragraphs)
    assert "EP 点" in joined or "EP" in joined
    assert "拓扑保护" in joined
    # 表格：3 行 2 列
    assert len(s2.tables) == 1
    assert len(s2.tables[0]) == 3
    assert s2.tables[0][1] == ["频率", "40 kHz"]
    # 图片：1 张
    assert len(s2.images) == 1
    # 备注
    assert "实验证据" in s2.notes


def test_read_pptx_export_images(sample_pptx: Path, tmp_path: Path) -> None:
    out = tmp_path / "imgs"
    slides = pptx_io.read_pptx(sample_pptx, export_images_to=out)
    imgs = [im for s in slides for im in s.images if im.get("path")]
    assert imgs, "应导出至少一张图片"
    assert Path(imgs[0]["path"]).exists()


def test_slides_to_markdown(sample_pptx: Path) -> None:
    slides = pptx_io.read_pptx(sample_pptx)
    md = pptx_io.slides_to_markdown(slides, source_name=sample_pptx.name)
    assert "第 1 页" in md and "第 2 页" in md
    assert "演讲者备注" in md
    assert "| 参数 | 值 |" in md  # 表格已渲染为 markdown
    assert "拓扑保护" in md


def test_pptx_to_markdown_one_shot(sample_pptx: Path) -> None:
    md = pptx_io.pptx_to_markdown(sample_pptx)
    assert "增益超表面" in md
    assert len(md) > 50


def test_extract_to_markdown_and_cache(sample_pptx: Path) -> None:
    res = extract.to_markdown(sample_pptx, force=True)
    assert res.backend in {"markitdown", "pptx_io"}
    assert res.char_count > 0
    assert res.cache_path is not None and res.cache_path.exists()
    # 二次调用命中缓存
    res2 = extract.to_markdown(sample_pptx)
    assert res2.from_cache is True


def test_reject_legacy_ppt(tmp_path: Path) -> None:
    fake = tmp_path / "old.ppt"
    fake.write_bytes(b"not a real ppt")
    with pytest.raises(ValueError):
        pptx_io.read_pptx(fake)
